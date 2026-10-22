"""Tables for a paper, and their export to CSV, Word, LaTeX and Markdown.

Each table has two forms. ``data`` holds the numbers at full precision, one
value per cell, and is what the CSV holds, so a reader can recompute from it.
``display`` holds the same numbers formatted for print, and is what goes into
the Word, LaTeX and Markdown files.
"""

from __future__ import annotations

import math
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pandas as pd


@dataclass
class Table:
    id: str
    title: str
    display: pd.DataFrame
    data: pd.DataFrame
    #: Lines printed under the table: abbreviations, tests, what n means.
    notes: list[str] = field(default_factory=list)
    #: Filled in when the table is numbered for the paper: "Table 2".
    number: str = ""

    @property
    def caption(self) -> str:
        head = f"{self.number}. " if self.number else ""
        return f"{head}{self.title}"


def decimals_for(values: pd.Series | np.ndarray, fixed: int | None = None) -> int:
    """Decimal places that show a variable's spread to about three significant figures."""
    if fixed is not None:
        return fixed
    numbers = pd.to_numeric(pd.Series(values), errors="coerce").dropna()
    if numbers.empty:
        return 2
    spread = float(numbers.std(ddof=1)) if len(numbers) > 1 else 0.0
    scale = spread if spread > 0 else float(numbers.abs().max())
    if not math.isfinite(scale) or scale == 0:
        return 0 if (numbers == numbers.round()).all() else 2
    return int(min(4, max(0, 2 - math.floor(math.log10(scale)))))


def number(value, decimals: int = 2) -> str:
    """A number for print with a true minus sign, or an en dash when there is none."""
    if value is None or (isinstance(value, float) and not math.isfinite(value)) or pd.isna(value):
        return "–"
    text = f"{float(value):,.{decimals}f}"
    if decimals and float(text.replace(",", "")) == 0:
        text = f"{0:.{decimals}f}"
    return text.replace("-", "−")


def p_value(value) -> str:
    """A p-value for print: three decimals, or "<0.001"."""
    if value is None or pd.isna(value):
        return "–"
    if value < 0.001:
        return "<0.001"
    return f"{value:.3f}"


def interval(low, high, decimals: int = 2) -> str:
    return f"{number(low, decimals)} to {number(high, decimals)}"


def write(table: Table, folder: Path, formats: tuple[str, ...]) -> list[Path]:
    """Write ``table`` in each of ``formats`` and return the paths, CSV first."""
    folder.mkdir(parents=True, exist_ok=True)
    written = []
    for kind in ("csv", "docx", "tex", "md"):
        if kind not in formats and kind != "csv":
            continue
        path = folder / f"{table.id}.{kind}"
        {"csv": _csv, "docx": _docx, "tex": _tex, "md": _markdown}[kind](table, path)
        written.append(path)
    return written


def _csv(table: Table, path: Path) -> None:
    table.data.to_csv(path, index=False, lineterminator="\n")


def markdown(table: Table) -> str:
    frame = table.display

    def cell(value) -> str:
        return str(value).replace("|", "\\|").replace("\n", " ")

    lines = [
        f"**{table.caption}**",
        "",
        "| " + " | ".join(cell(c) for c in frame.columns) + " |",
        "|" + "|".join("---" for _ in frame.columns) + "|",
    ]
    lines += ["| " + " | ".join(cell(value) for value in row) + " |" for row in frame.itertuples(index=False)]
    if table.notes:
        lines += [""] + [f"*{note}*" if index == 0 else note for index, note in enumerate(table.notes)]
    return "\n".join(lines) + "\n"


def _markdown(table: Table, path: Path) -> None:
    path.write_text(markdown(table), encoding="utf-8")


_LATEX_SPECIAL = {
    "\\": r"\textbackslash{}",
    "&": r"\&",
    "%": r"\%",
    "$": r"\$",
    "#": r"\#",
    "_": r"\_",
    "{": r"\{",
    "}": r"\}",
    "~": r"\textasciitilde{}",
    "^": r"\textasciicircum{}",
    "−": "$-$",
    "<": r"\textless{}",
    ">": r"\textgreater{}",
    "±": r"$\pm$",
    "–": "--",
    "×": r"$\times$",
    "²": r"\textsuperscript{2}",
    "≥": r"$\geq$",
    "≤": r"$\leq$",
}


def latex_escape(text) -> str:
    return "".join(_LATEX_SPECIAL.get(character, character) for character in str(text))


def latex(table: Table) -> str:
    """A booktabs table: needs ``\\usepackage{booktabs}`` in the preamble."""
    frame = table.display
    align = "l" + "r" * (len(frame.columns) - 1)
    lines = [
        "\\begin{table}[ht]",
        "\\centering",
        f"\\caption{{{latex_escape(table.title)}}}",
        f"\\label{{tab:{table.id}}}",
        f"\\begin{{tabular}}{{{align}}}",
        "\\toprule",
        " & ".join(latex_escape(column) for column in frame.columns) + " \\\\",
        "\\midrule",
    ]
    lines += [
        " & ".join(latex_escape(value) for value in row) + " \\\\" for row in frame.itertuples(index=False)
    ]
    lines += ["\\bottomrule", "\\end{tabular}"]
    if table.notes:
        lines += [
            "\\par\\smallskip",
            "{\\footnotesize " + " ".join(latex_escape(note) for note in table.notes) + "}",
        ]
    lines.append("\\end{table}")
    return "\n".join(lines) + "\n"


def _tex(table: Table, path: Path) -> None:
    path.write_text(latex(table), encoding="utf-8")


def _docx(table: Table, path: Path) -> None:
    write_docx([table], path)


def write_docx(tables: list[Table], path: Path) -> None:
    """Write tables to one Word file, each with its caption above and notes below, all editable."""
    document = docx_document()
    for position, table in enumerate(tables):
        if position:
            document.add_paragraph()
        add_docx_table(document, table)
    save_docx(document, path)


def docx_document():
    """A new Word document in the tables' font, with metadata fixed so rebuilds match."""
    from docx import Document  # noqa: PLC0415
    from docx.shared import Pt  # noqa: PLC0415

    document = Document()
    _stable_properties(document)
    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)
    return document


def add_docx_table(document, table: Table) -> None:
    """Add ``table`` to a Word document: caption above, an editable grid, notes below."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: PLC0415
    from docx.shared import Pt  # noqa: PLC0415

    caption = document.add_paragraph()
    run = caption.add_run(f"{table.number}. " if table.number else "")
    run.bold = True
    caption.add_run(table.title)
    frame = table.display
    grid = document.add_table(rows=1, cols=len(frame.columns))
    grid.style = "Table Grid"
    for cell, column in zip(grid.rows[0].cells, frame.columns, strict=True):
        cell.text = ""
        cell.paragraphs[0].add_run(str(column)).bold = True
    for row in frame.itertuples(index=False):
        cells = grid.add_row().cells
        for index, (cell, value) in enumerate(zip(cells, row, strict=True)):
            cell.text = str(value)
            if index:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for note in table.notes:
        paragraph = document.add_paragraph()
        paragraph.add_run(note).font.size = Pt(8)


#: The date stamped on every part of a Word file, so the same content gives the same bytes.
DOCX_TIMESTAMP = (2000, 1, 1, 0, 0, 0)


def save_docx(document, path: Path) -> None:
    """Save a python-docx ``Document`` so the same content gives the same file at every build."""
    _stable_properties(document)
    document.save(path)
    stable_docx(path)


def stable_docx(path: Path) -> None:
    """Rewrite a Word file with a fixed date on each part of its zip archive.

    python-docx stamps every part with the time of saving, so two builds of the
    same table would otherwise differ in bytes and show up in ``git diff``.
    """
    import io  # noqa: PLC0415
    import zipfile  # noqa: PLC0415

    with zipfile.ZipFile(path) as source:
        parts = [(info.filename, info.external_attr, source.read(info)) for info in source.infolist()]
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for name, attributes, content in parts:
            info = zipfile.ZipInfo(name, date_time=DOCX_TIMESTAMP)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = attributes
            target.writestr(info, content)
    path.write_bytes(buffer.getvalue())


def _stable_properties(document) -> None:
    """Fix the Word file's metadata so the same table gives the same file on every build."""
    from datetime import datetime  # noqa: PLC0415

    properties = document.core_properties
    fixed = datetime(2000, 1, 1)
    properties.created = fixed
    properties.modified = fixed
    properties.last_printed = fixed
    properties.author = ""
    properties.last_modified_by = ""
    properties.revision = 1
