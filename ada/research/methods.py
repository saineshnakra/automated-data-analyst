"""The methods draft, a results draft and the supplement, written from what the build ran.

Every sentence comes from the build itself: the processing record, each
analysis result, the figure assets and the installed package versions. A step
that did not run is never described, so the draft cannot claim an analysis the
data did not have. It is a draft: the authors add what ADA cannot know, such as
how the data were collected, and edit the wording.

Each document is built once as a list of blocks and written from that list to
Markdown and to Word, so the two files always say the same thing.
"""

from __future__ import annotations

import math
import os
import platform
import re
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd

from ..file_io import EXCEL_SUFFIXES
from . import provenance
from .audit import EXTREME_Z
from .descriptives import grouping
from .processing import levels_of
from .provenance import Asset
from .tables import Table, add_docx_table, docx_document, markdown, number, p_value, save_docx

#: Every part of study.yml. The drafts describe the whole study, so any change to it can change them.
WHOLE_STUDY = (
    "study",
    "data",
    "design",
    "variables",
    "processing",
    "decisions",
    "analyses",
    "figures",
    "descriptives",
    "output",
    "sharing",
    "custom",
)

#: Each multiple-comparison method: its name in a sentence, its short name, and its reference.
ADJUSTMENTS = {
    "holm": (
        "Holm's step-down method",
        "Holm",
        "Holm, S. (1979). A simple sequentially rejective multiple test procedure. "
        "Scandinavian Journal of Statistics 6, 65–70.",
    ),
    "bonferroni": (
        "the Bonferroni correction",
        "Bonferroni",
        "Dunn, O. J. (1961). Multiple comparisons among means. Journal of the American Statistical "
        "Association 56, 52–64. https://doi.org/10.1080/01621459.1961.10482090",
    ),
    "fdr_bh": (
        "the Benjamini-Hochberg false discovery rate procedure",
        "Benjamini-Hochberg",
        "Benjamini, Y. and Hochberg, Y. (1995). Controlling the false discovery rate: a practical and "
        "powerful approach to multiple testing. Journal of the Royal Statistical Society B 57, 289–300. "
        "https://doi.org/10.1111/j.2517-6161.1995.tb02031.x",
    ),
}

#: The libraries a build uses, in the order the draft names them, with how to cite each (docs/credits.md).
LIBRARIES = (
    (
        "pandas",
        "pandas",
        "McKinney, W. (2010). Data structures for statistical computing in Python. Proceedings of the "
        "9th Python in Science Conference, 56–61. https://doi.org/10.25080/Majora-92bf1922-00a",
    ),
    (
        "numpy",
        "NumPy",
        "Harris, C. R., Millman, K. J., van der Walt, S. J. et al. (2020). Array programming with NumPy. "
        "Nature 585, 357–362. https://doi.org/10.1038/s41586-020-2649-2",
    ),
    (
        "scipy",
        "SciPy",
        "Virtanen, P., Gommers, R., Oliphant, T. E. et al. (2020). SciPy 1.0: fundamental algorithms for "
        "scientific computing in Python. Nature Methods 17, 261–272. https://doi.org/10.1038/s41592-019-0686-2",
    ),
    (
        "statsmodels",
        "statsmodels",
        "Seabold, S. and Perktold, J. (2010). statsmodels: Econometric and statistical modeling with "
        "Python. Proceedings of the 9th Python in Science Conference, 92–96. "
        "https://doi.org/10.25080/Majora-92bf1922-011",
    ),
    (
        "matplotlib",
        "Matplotlib",
        "Hunter, J. D. (2007). Matplotlib: a 2D graphics environment. Computing in Science & Engineering "
        "9, 90–95. https://doi.org/10.1109/MCSE.2007.55",
    ),
    ("python-docx", "python-docx", ""),
    ("PyYAML", "PyYAML", ""),
)
PYTHON_REFERENCE = "Van Rossum, G. and Drake, F. L. (2009). Python 3 Reference Manual. CreateSpace."
ADA_REFERENCE = (
    "Nakra, S. ADA: Automated Data Analyst (automated-data-analyst), version {version}. "
    "https://github.com/saineshnakra/automated-data-analyst"
)
GENERATED = (
    "Generated from the steps this build executed: the processing record, each analysis result (its test, "
    "rows used and left out, planned status and adjustment), the figure assets and the installed package "
    "versions. Nothing is written about a step that did not run."
)

_WORDS = ("no", "one", "two", "three", "four", "five", "six", "seven", "eight", "nine")
_INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
#: Names that take "the" in front of a test: "the Mann-Whitney U test".
_NAMED_TESTS = {
    "mann",
    "kruskal",
    "wilcoxon",
    "fisher",
    "friedman",
    "mcnemar",
    "pearson",
    "spearman",
    "kendall",
    "student",
    "tukey",
    "kolmogorov",
    "shapiro",
    "levene",
}


@dataclass
class Block:
    """One piece of a document: a heading, paragraph, list, table, code listing or picture."""

    kind: str
    text: str = ""
    level: int = 2
    items: tuple[str, ...] = ()
    table: Table | None = None
    image: Path | None = None


@dataclass
class Draft:
    """A document as blocks, written the same way to Markdown and to Word."""

    title: str
    blocks: list[Block] = field(default_factory=list)

    def heading(self, text: str, level: int = 2) -> None:
        self.blocks.append(Block("heading", text, level))

    def paragraph(self, *sentences: str) -> None:
        text = " ".join(sentence for sentence in sentences if sentence)
        if text:
            self.blocks.append(Block("paragraph", text))

    def note(self, text: str) -> None:
        self.blocks.append(Block("note", text))

    def bullets(self, items) -> None:
        self.blocks.append(Block("bullets", items=tuple(items)))

    def numbered(self, items) -> None:
        self.blocks.append(Block("numbered", items=tuple(items)))

    def add_table(self, table: Table) -> None:
        self.blocks.append(Block("table", table=table))

    def code(self, text: str) -> None:
        self.blocks.append(Block("code", text.rstrip("\n")))

    def picture(self, path: Path, caption: str) -> None:
        self.blocks.append(Block("image", caption, image=path))


class References:
    """References numbered in the order the text first cites them."""

    def __init__(self) -> None:
        self.items: list[str] = []

    def cite(self, *references: str) -> str:
        numbers = []
        for reference in references:
            if not reference:
                continue
            if reference not in self.items:
                self.items.append(reference)
            numbers.append(self.items.index(reference) + 1)
        return "[" + ", ".join(map(str, sorted(set(numbers)))) + "]" if numbers else ""


def run(result) -> None:
    """Write the methods draft, the results draft and the supplement into ``result.out``."""
    folder = result.folder("methods")
    methods, references = methods_draft(result)
    methods_files = write(methods, folder / "methods", folder)
    results_path = folder / "results.md"
    results_path.write_text(to_markdown(results_draft(result), folder), encoding="utf-8")
    supplement_folder = result.folder("supplement")
    supplement_files = write(supplement_draft(result), supplement_folder / "supplement", supplement_folder)

    entries = _analyses(result)
    done = [found for _, found in entries if found is not None and found.supported]
    rows = result.rows()
    result.assets.append(
        Asset(
            id="methods-draft",
            kind="text",
            title="Methods draft",
            files=methods_files,
            depends=WHOLE_STUDY,
            calculation=GENERATED,
            references=tuple(references.items),
            code="ada.research.methods.methods_draft",
            rows=rows,
            numbers={
                "analyses": len(entries),
                "planned": sum(1 for spec, _ in entries if spec.planned),
                "analysed": len(done),
                "references": len(references.items),
            },
            section="record",
        )
    )
    result.assets.append(
        Asset(
            id="results-draft",
            kind="text",
            title="Results draft",
            files=[results_path],
            depends=WHOLE_STUDY,
            calculation=GENERATED
            + " Estimates, intervals and p-values are copied from the analysis results.",
            code="ada.research.methods.results_draft",
            rows=rows,
            numbers={"analyses_reported": len(done)},
            section="record",
        )
    )
    supplementary = [table for table in result.tables if table.number.startswith("Table S")]
    result.assets.append(
        Asset(
            id="supplement",
            kind="text",
            title="Supplementary material",
            files=supplement_files,
            depends=WHOLE_STUDY,
            calculation="Collected from the steps this build executed: the supplementary tables and figures, "
            "the data audit's items to check, the processing record, each model's output and the data "
            "dictionary.",
            code="ada.research.methods.supplement_draft",
            rows=rows,
            numbers={"tables": len(supplementary), "figures": len(_figures(result, "supplement"))},
            number="Supplement",
            section="supplement",
        )
    )


# The methods draft.


def methods_draft(result) -> tuple[Draft, References]:
    """The Methods section as the build ran it, with its numbered references."""
    references = References()
    draft = Draft("Methods (draft)")
    draft.note(
        "Draft for the authors to edit. ADA wrote it from what this build ran: the processing record, the "
        "analysis results, the figures and the software versions. Check every sentence and add what ADA "
        "cannot know, such as how and why the data were collected and any ethics approval. `ada build` "
        "rewrites this file, so copy the text into the manuscript before editing it."
    )
    _data_section(draft, result)
    _processing_section(draft, result, references)
    _analysis_section(draft, result, references)
    _figures_section(draft, result, references)
    _software_section(draft, result, references)
    draft.heading("References")
    if references.items:
        draft.numbered(references.items)
    else:
        draft.paragraph("No references were cited.")
    return draft, references


def _data_section(draft: Draft, result) -> None:
    study, processed = result.study, result.processed
    original = processed.original
    read_step = _record(processed, "read")
    rows_in_file = read_step.rows_before if read_step else len(original)
    source = study.data_path + (f" (sheet {study.data_sheet})" if study.data_sheet else "")
    checked = ", checked against the value recorded in study.yml" if study.data_sha256 else ""
    draft.heading("Data")
    draft.paragraph(
        f"The data were read from {source}, which held {_n(rows_in_file, 'row')} and "
        f"{_count(len(original.columns) - 1, 'column')} (SHA-256 checksum beginning "
        f"{processed.sha256[:12]}{checked}; the full value is in provenance.json).",
        _observation(study.observation),
    )
    draft.paragraph(*_design_sentences(study, processed.data))
    draft.paragraph(*_variable_sentences(study, processed))
    draft.paragraph(*_missing_sentences(study, processed))


def _observation(text: str) -> str:
    text = text.strip().rstrip(".")
    if not text:
        return ""
    if text.lower().startswith(("one row per", "a row per", "row per")):
        return f"The file held {text}."
    return f"Each row represented {text}."


def _design_sentences(study, data: pd.DataFrame) -> list[str]:
    """The design with its counts in the analysed rows: subjects, clusters, time points and groups."""
    subject = study.subject if study.subject in data else None
    cluster = study.cluster if study.cluster in data else None
    time = study.time if study.time in data else None
    rows = len(data)
    sentences = []
    if study.design == "independent":
        sentences.append("The design was independent: each row came from a different subject.")
    elif study.design == "paired":
        sentences.append(_paired_sentence(study, data, subject))
    elif study.design == "repeated":
        if subject:
            per = data[subject].value_counts()
            sentences.append(
                f"The design was repeated measures: the {_n(rows, 'row')} analysed came from "
                f"{_n(per.size, 'subject')}, identified by {subject}, with {_spread(per)} rows each."
            )
        else:
            sentences.append("The design was repeated measures.")
    elif study.design == "clustered":
        if cluster:
            per = data[cluster].value_counts()
            from_subjects = f" and {_n(data[subject].nunique(), 'subject')} ({subject})" if subject else ""
            sentences.append(
                f"The design was clustered: the {_n(rows, 'row')} analysed came from "
                f"{_n(per.size, 'cluster')}, identified by {cluster}{from_subjects}, with {_spread(per)} "
                "rows per cluster."
            )
        else:
            sentences.append("The design was clustered.")
    if time:
        moments = _moments(data, time)
        listed = (
            _join([_value(moment) for moment in moments])
            if len(moments) <= 10
            else f"{_value(moments[0])} to {_value(moments[-1])}"
        )
        sentence = f"Measurements were made at {_count(len(moments), 'time point')} ({time} {listed})"
        if subject:
            per_time = data.groupby(time, observed=True)[subject].nunique()
            sentence += f", with {_spread(per_time)} subjects at each"
        sentences.append(sentence + ".")
    for group in study.with_role("group"):
        if group not in data or (study.design == "paired" and group == _condition(study, data)):
            continue
        levels = levels_of(data, group)
        if not levels:
            continue
        labels = data[group].astype("string")
        between = subject and (data.groupby(subject, observed=True)[group].nunique() <= 1).all()
        if between:
            counts = [int(data.loc[labels == level, subject].nunique()) for level in levels]
            noun = "subject"
        else:
            counts = [int((labels == level).sum()) for level in levels]
            noun = "row"
            if subject:
                sentences.append(
                    f"Each subject was measured under the {_count(len(levels), 'level')} of {group} "
                    f"({_join(levels)})."
                )
        first = f"{_n(counts[0], noun)} {_was(counts[0])} in {levels[0]}"
        rest = [f"{count:,} in {level}" for level, count in zip(levels[1:], counts[1:], strict=True)]
        sentences.append(f"By {group}, {_join([first, *rest])}.")
    if study.design_notes:
        sentences.append(_sentence(study.design_notes))
    return sentences


def _condition(study, data: pd.DataFrame) -> str | None:
    return next((name for name in study.with_role("group") if name in data), None)


def _paired_sentence(study, data: pd.DataFrame, subject: str | None) -> str:
    condition = _condition(study, data)
    if not subject or not condition:
        return "The design was paired."
    levels = levels_of(data, condition)
    seen = data.groupby(subject, observed=True)[condition].nunique()
    complete = int((seen == len(levels)).sum())
    partial = int(seen.size - complete)
    every = "both conditions" if len(levels) == 2 else "every condition"
    text = (
        f"The design was paired: each subject, identified by {subject}, was measured under the "
        f"{_count(len(levels), 'condition')} of {condition} ({_join(levels)}). Of the "
        f"{_n(seen.size, 'subject')}, {complete:,} had {every}"
    )
    if partial:
        text += f"; {partial:,} had {'only one' if len(levels) == 2 else 'fewer'} and could not be paired"
    return text + "."


def _variable_sentences(study, processed) -> list[str]:
    data = processed.data
    used = _used_variables(study)
    sentences = []

    def described(name: str) -> str:
        variable = study.variable(name)
        text = variable.display
        if variable.description:
            text += f" ({variable.description.rstrip('.')})"
        if (
            variable.type == "numeric"
            and not variable.unit
            and variable.role in ("outcome", "covariate", "time")
        ):
            text += " [unit not recorded in study.yml]"
        return text

    outcomes = study.with_role("outcome")
    if outcomes:
        verb = "outcome was" if len(outcomes) == 1 else "outcomes were"
        sentences.append(f"The {verb} {_join([described(name) for name in outcomes])}.")
    for name in study.with_role("group"):
        variable = study.variable(name)
        levels = [str(level) for level in variable.levels] or (levels_of(data, name) if name in data else [])
        if levels:
            sentences.append(
                f"{variable.display} defined the groups ({_join(levels)}), with {levels[0]} as the "
                "reference level."
            )
        else:
            sentences.append(f"{variable.display} defined the groups.")
    for name in study.with_role("time"):
        sentences.append(f"{described(name)} was the time variable.")
    covariates = study.with_role("covariate")
    if covariates:
        verb = "covariate was" if len(covariates) == 1 else "covariates were"
        sentences.append(f"The {verb} {_join([described(name) for name in covariates])}.")
    identifiers = [
        name for name in study.with_role("identifier") if name not in (study.subject, study.cluster)
    ]
    if identifiers:
        sentences.append(f"Rows were identified by {_join(identifiers)}.")
    unused = [name for name in study.with_role("other") if name not in used]
    if unused:
        sentences.append(f"{_join(unused)} {_was(len(unused))} not used in any analysis.")
    if processed.undescribed:
        count = len(processed.undescribed)
        sentences.append(
            f"The file also held {_count(count, 'column')} not described in study.yml "
            f"({_join(processed.undescribed)}); {'it' if count == 1 else 'they'} {_was(count)} not used."
        )
    sentences.append(
        "Each variable's type, unit, allowed values and observed range are listed in the data dictionary "
        "(supplementary material)."
    )
    return sentences


def _codes(study) -> str:
    """The missing-value codes in study.yml: "-99 in every column; 999 in Score"."""
    parts = (
        [f"{_join([str(code) for code in study.missing_codes])} in every column"]
        if study.missing_codes
        else []
    )
    parts += [
        f"{_join([str(code) for code in variable.missing_codes])} in {name}"
        for name, variable in study.variables.items()
        if variable.missing_codes
    ]
    return "; ".join(parts)


def _used_variables(study) -> set[str]:
    used = {name for name in (study.subject, study.cluster, study.time) if name}
    for spec in study.analyses:
        used.update(name for name in (spec.outcome, spec.compare, spec.relate, spec.change, spec.by) if name)
        used.update(spec.covariates)
    return used


def _missing_sentences(study, processed) -> list[str]:
    data = processed.data
    sentences = []
    found = processed.missing_codes
    total = sum(sum(counts.values()) for counts in found.values())
    codes = _codes(study)
    if codes:
        where = "; ".join(f"{sum(counts.values()):,} in {name}" for name, counts in found.items())
        sentences.append(f"Missing values were coded as {codes}.")
        sentences.append(
            f"Using these codes, {_n(total, 'value')} {_was(total)} read as missing ({where})."
            if total
            else "No such codes occurred in the data."
        )
    else:
        sentences.append("No missing-value codes were declared; empty cells were read as missing.")
    read_step = _record(processed, "read")
    if read_step and read_step.values_changed:
        sentences.append(
            f"When the file was read, {_n(read_step.values_changed, 'further value')} "
            f"{_was(read_step.values_changed)} recognised as missing (processing step 1)."
        )
    for name, (count, examples) in processed.unreadable.items():
        sentences.append(
            f"In {name}, {_n(count, 'value')} could not be read as {study.variable(name).type} and "
            f"{_was(count)} treated as missing (for example, {', '.join(examples)})."
        )
    used = _used_variables(study)
    columns = [name for name in study.variables if name in data and name in used]
    gaps = [(name, int(data[name].isna().sum())) for name in columns]
    gaps = [(name, count) for name, count in gaps if count]
    if gaps:
        listed = _join([f"{name} ({count:,} of {len(data):,} rows)" for name, count in gaps])
        sentences.append(f"After processing, values were missing for {listed}.")
    elif columns:
        sentences.append("After processing, no value was missing in any analysed variable.")
    return sentences


def _processing_section(draft: Draft, result, references: References) -> None:
    study, processed = result.study, result.processed
    draft.heading("Data processing")
    draft.paragraph(
        "The data were processed in the following steps, each applied to the rows left by the one before. "
        "The step numbers are those of the processing record (processing/processing-record.md)."
    )
    steps = iter(study.steps)
    items = []
    for record in processed.records:
        step = next(steps, None) if record.kind in ("exclude", "transform", "recode") else None
        items.append(f"**Step {record.number}.** " + _step_text(record, step, study, processed))
    draft.bullets(items)

    read_step = _record(processed, "read")
    rows_read = read_step.rows_before if read_step else len(processed.original)
    excluded = len(processed.excluded) if processed.excluded is not None else 0
    if excluded == 1:
        kept = (
            ", and the excluded row was kept with the step that removed it and its line in the original file "
            "(processing/excluded-rows.csv)"
        )
    elif excluded:
        kept = (
            f", and the {excluded:,} excluded rows were kept, each with the step that removed it and its "
            "line in the original file (processing/excluded-rows.csv)"
        )
    else:
        kept = ", and no rows were excluded"
    draft.paragraph(
        f"After processing, {len(processed.data):,} of the {rows_read:,} rows read remained for analysis."
        if len(processed.data) != rows_read
        else f"All {rows_read:,} rows read remained for analysis.",
        f"The original file was not changed{kept}.",
    )

    audit_asset = _asset(result, "data-audit")
    cite = references.cite(*(audit_asset.references if audit_asset else ()))
    findings = getattr(result, "findings", []) or []
    extremes = [finding for finding in findings if finding.area == "extremes"]
    checks = sum(1 for finding in findings if finding.status == "check")
    marked = f"marked {_count(checks, 'item')} for checking" if checks else "marked nothing for checking"
    sentences = [
        "Before analysis, the data were checked for missing, unreadable, out-of-range and duplicated values "
        f"and for consistency with the stated design; the data audit {marked} (audit/data-audit.md, "
        "summarised in the supplementary material).",
        f"Extreme values in the outcomes and covariates were defined as those with a modified z-score above "
        f"{EXTREME_Z:g} in absolute value, that is, more than {EXTREME_Z:g} robust standard deviations from "
        "the median" + (f" {cite}" if cite else "") + ".",
    ]
    if extremes:
        count = sum(finding.count for finding in extremes)
        where = "; ".join(f"{finding.count:,} in {finding.column}" for finding in extremes)
        sentences.append(f"There {'was' if count == 1 else 'were'} {_n(count, 'such value')} ({where}).")
    else:
        sentences.append("There were none.")
    sentences.append("No value was removed for being extreme.")
    draft.paragraph(*sentences)


def _step_text(record, step, study, processed) -> str:
    """One processing step in past-tense prose, with the rows or values it affected and its reason."""
    if record.kind == "read":
        source = study.data_path + (f" (sheet {study.data_sheet})" if study.data_sheet else "")
        dropped = record.rows_before - record.rows_after
        text = (
            f"{source} was read: {_n(record.rows_before, 'row')} read and {record.rows_after:,} kept"
            + (f" ({_n(dropped, 'empty row')} dropped)" if dropped else "")
            + f"; {_n(record.values_changed, 'value')} set to missing by the reader. "
            + _sentence(record.reason)
        )
        notes = " ".join(_sentence(note) for note in processed.reading_notes)
        return text + (f" {notes}" if notes else "")
    if record.kind == "missing-codes":
        where = f", in {_join(record.columns)}" if record.columns else ""
        return (
            f"The missing-value codes listed in study.yml ({_codes(study)}) were read as missing: "
            f"{_n(record.values_changed, 'value')} set to missing{where}."
        )
    if record.kind == "types":
        examples = "; ".join(
            f"{name}: {', '.join(found[1])}" for name, found in processed.unreadable.items() if found[1]
        )
        return (
            f"Values that could not be read as their declared type were set to missing: "
            f"{_n(record.values_changed, 'value')}, in {_join(record.columns)}"
            + (f" (for example, {examples})" if examples else "")
            + f". {_sentence(record.reason)}"
        )
    reason = f" Reason: {_sentence(record.reason)}" if record.reason else ""
    if record.kind == "exclude":
        text = _exclusion(step) if step else record.description
        text += f" ({_n(record.rows_removed, 'row')} removed, leaving {record.rows_after:,})."
        if record.removed_ids and not (step and step.ids):
            column = (step.id_column if step else None) or study.subject or "identifier"
            shown = list(record.removed_ids[:10])
            more = len(record.removed_ids) - len(shown)
            text += f" The removed rows had {column} {_join(shown + ([f'{more} others'] if more else []))}."
        return text + reason
    if record.kind == "transform":
        text = _transformation(step) if step else record.description
        return text + f" ({_n(record.values_changed, 'value')} computed)." + reason
    text = _recoding(step) if step else record.description
    return text + f" ({_n(record.values_changed, 'value')} changed)." + reason


def _exclusion(step) -> str:
    parts = [_condition_text(condition) for condition in step.where]
    if step.ids:
        parts.append(f"{step.id_column} was {_join([str(value) for value in step.ids], 'or')}")
    return "Rows in which " + " and ".join(parts) + " were excluded"


def _condition_text(condition) -> str:
    """A ``where`` condition in the past tense: "Viability (%) was below 0"."""
    column, value = condition.column, condition.value
    if condition.test == "missing":
        return f"{column} was missing" if value in (None, True) else f"{column} had a value"
    if condition.test in ("in", "not_in"):
        values = [str(item) for item in (value if isinstance(value, list | tuple) else [value])]
        return f"{column} {'was' if condition.test == 'in' else 'was not'} {_join(values, 'or')}"
    words = {
        "below": "was below",
        "above": "was above",
        "at_most": "was at most",
        "at_least": "was at least",
        "equals": "was",
        "not_equals": "was not",
    }
    return f"{column} {words[condition.test]} {value}"


def _transformation(step) -> str:
    column = step.column
    operation = {
        "log": f"the natural logarithm of {column}",
        "log10": f"the base-10 logarithm of {column}",
        "log2": f"the base-2 logarithm of {column}",
        "sqrt": f"the square root of {column}",
        "square": f"the square of {column}",
        "reciprocal": f"the reciprocal of {column}",
        "zscore": f"the z-score of {column} (its mean subtracted and divided by its SD, both over the rows "
        "at this step)",
        "center": f"{column} minus its mean over the rows at this step",
        "scale": f"{column} multiplied by {step.factor:g}" if step.factor is not None else f"scaled {column}",
    }.get(step.apply, f"{step.apply}({column})")
    return f"{step.to or column} was set to {operation}"


def _recoding(step) -> str:
    pairs = ", ".join(f"{old} → {new}" for old, new in step.mapping.items())
    into = f" into {step.to}" if step.to and step.to != step.column else ""
    return f"{step.column} was recoded{into}: {pairs}"


def _analysis_section(draft: Draft, result, references: References) -> None:
    study = result.study
    decisions = study.decisions
    entries = _analyses(result)
    names = _names(entries)
    draft.heading("Statistical analysis")
    draft.paragraph(*_descriptive_sentences(result))
    if not entries:
        draft.paragraph(
            "No statistical analyses were listed in study.yml and none were run; the tables and figures "
            "describe the data only."
        )
        _custom_paragraph(draft, result)
        return
    planned = [spec for spec, _ in entries if spec.planned]
    exploratory = [spec for spec, _ in entries if not spec.planned]
    sentences = [
        f"The analysis plan (study.yml) listed {_count(len(entries), 'analysis', 'analyses')}, and no other "
        "analyses were run."
    ]
    if len(entries) == 1:
        sentences.append(
            "It was declared there as planned before the data were seen."
            if planned
            else "It was exploratory: it was not declared as planned before the data were seen."
        )
    elif not exploratory:
        sentences.append("All were declared there as planned before the data were seen.")
    elif not planned:
        sentences.append("None was declared as planned before the data were seen, so all are exploratory.")
    else:
        sentences.append(
            f"{_cap(_count(len(planned), 'analysis', 'analyses'))} {_was(len(planned))} declared there as "
            f"planned before the data were seen ({_join([names[spec.id] for spec in planned])}); the other "
            + (f"{_count(len(exploratory), 'analysis', 'analyses')} were" if len(exploratory) > 1 else "was")
            + f" exploratory ({_join([names[spec.id] for spec in exploratory])})."
        )
    if exploratory:
        sentences.append(
            "Exploratory results are labelled as such wherever they are reported and should be read as "
            "suggesting hypotheses, not testing them."
        )
    sentences.append(
        "Missing data were handled by complete-case analysis: each analysis used the rows with every value "
        "it needed, and no values were imputed."
    )
    draft.paragraph(*sentences)

    for spec, found in entries:
        draft.paragraph(*_analysis_sentences(spec, found, study, references))

    done = [found for _, found in entries if found is not None and found.supported]
    if done:
        families: dict[str, list] = {}
        for found in sorted(done, key=lambda found: not found.planned):
            families.setdefault(found.family or found.spec.id, []).append(found.spec)
        listed = _join(
            [
                f"“{family}” ({_join([names[spec.id] for spec in specs])})"
                for family, specs in families.items()
            ]
        )
        method = decisions.multiple_comparisons
        if method in ADJUSTMENTS:
            name, _, reference = ADJUSTMENTS[method]
            adjusted = (
                f"P-values were adjusted for multiple comparisons with {name} {references.cite(reference)} "
                f"within each family of analyses named in study.yml: {listed}."
            )
        else:
            adjusted = "P-values were not adjusted for multiple comparisons."
        draft.paragraph(
            adjusted,
            f"The significance level was α = {decisions.alpha:g}, and all confidence intervals were computed "
            f"at the {_percent(decisions.confidence)} level.",
        )
    _custom_paragraph(draft, result)


def _descriptive_sentences(result) -> list[str]:
    """What the descriptive tables report, for the tables this build made."""
    study = result.study
    sentences = []
    table = _asset(result, "descriptives")
    if table is not None:
        by = grouping(study)
        sentences.append(
            f"{table.number or 'The descriptive table'} summarises each variable"
            + (f" overall and by {by}" if by else "")
            + ": numeric variables as mean and standard deviation (n − 1 denominator), median and "
            "interquartile range (quartiles by linear interpolation) and range; categorical variables as "
            "counts and percentages of the rows with a value."
        )
    over_time = _asset(result, "outcomes-over-time")
    if over_time is not None and study.time:
        by = grouping(study)
        sentences.append(
            f"{over_time.number or 'A second table'} gives each outcome's mean, standard deviation and n at "
            f"each value of {study.time}" + (f", by {by}" if by else "") + "."
        )
    return sentences


def _analysis_sentences(spec, found, study, references: References) -> list[str]:
    status = "planned" if spec.planned else "exploratory"
    lead = f"**{_title(spec)}** (`{spec.id}`)."
    if found is None:
        return [
            lead,
            f"This {status} analysis was listed in study.yml, but the build produced no result for it.",
        ]
    if not found.supported:
        return [lead, f"This {status} analysis was not run.", _sentence(found.reason) if found.reason else ""]
    sentences = [lead]
    if spec.question:
        sentences.append(f"It addressed the question “{spec.question.strip()}”")
        if not spec.question.strip().endswith(("?", ".", "!")):
            sentences[-1] += "."
    test = _with_article(found.test) + (f" {references.cite(*found.references)}" if found.references else "")
    method = f" (`{found.method}`)" if found.method else ""
    groups = "; ".join(f"{level}, {_n(count, 'row')}" for level, count in found.groups.items())
    in_groups = f" ({groups})" if groups else ""
    if spec.kind == "compare":
        sentences.append(
            f"{spec.outcome} was compared between the levels of {spec.compare}{in_groups} with "
            f"{test}{method}."
        )
    elif spec.kind == "relate":
        sentences.append(
            f"The association between {spec.outcome} and {spec.relate} was estimated with {test}{method}."
        )
    else:
        by = f", and whether it differed between the levels of {spec.by}{in_groups}," if spec.by else ""
        sentences.append(
            f"The change in {spec.outcome} over {spec.change}{by} was estimated with {test}{method}."
        )
    if found.formula:
        sentences.append(f"The model was `{found.formula}`.")
    if spec.covariates:
        kind = "fixed-effect covariates" if len(spec.covariates) > 1 else "a fixed-effect covariate"
        sentences.append(f"{_join(spec.covariates)} entered the model as {kind}.")
    elif found.formula:
        sentences.append("No covariates were included.")
    units = []
    if found.subjects:
        units.append(_n(found.subjects, "subject"))
    if found.clusters:
        units.append(_n(found.clusters, "cluster"))
    source = f", from {_join(units)}," if units else ""
    if found.n_missing:
        sentences.append(
            f"The analysis used the {_n(found.n, 'row')}{source} that had every value it needed (complete "
            f"cases); {_n(found.n_missing, 'row')} {_was(found.n_missing)} left out because a value was "
            "missing."
        )
    else:
        sentences.append(
            f"The analysis used {_n(found.n, 'row')}{source.rstrip(',')}; none was left out for a missing "
            "value."
        )
    effect = found.effect.get("name") if found.effect else ""
    if effect:
        sentences.append(
            f"{effect} was reported as the standardised effect size, with its "
            f"{_percent(study.decisions.confidence)} confidence interval."
        )
    sentences.append(
        "This analysis was planned before the data were seen."
        if spec.planned
        else "This analysis was exploratory: it was not planned before the data were seen."
    )
    if found.diagnostics:
        sentences.append("Assumption checks: " + " ".join(_sentence(d) for d in found.diagnostics))
    for warning in found.warnings:
        sentences.append(f"Caveat: {_sentence(warning)}")
    return sentences


def _custom_paragraph(draft: Draft, result) -> None:
    """Assets made by the authors' own code: ADA ran it but cannot describe its methods."""
    custom = getattr(result, "custom", {}) or {}
    for script, digest in custom.items():
        made = [asset for asset in result.assets if asset.code.startswith(f"{script}:")]
        if not made:
            continue
        names = _join([f"{asset.number} ({asset.title})" if asset.number else asset.title for asset in made])
        draft.paragraph(
            f"{names} {_was(len(made))} made by the authors' own code in {script} (SHA-256 beginning "
            f"{digest[:12]}), run on the processed data. Describe its methods here; ADA does not."
        )


def _figures_section(draft: Draft, result, references: References) -> None:
    draft.heading("Figures")
    figures = _figures(result, "main") + _figures(result, "supplement")
    if not figures:
        draft.paragraph("No figures were drawn.")
        return
    style = getattr(result, "style", None)
    version = provenance.package_versions().get("matplotlib", "")
    library = next(entry for entry in LIBRARIES if entry[0] == "matplotlib")
    drawn = " ".join(
        part for part in ("Figures were drawn with Matplotlib", version, references.cite(library[2])) if part
    )
    if style is not None:
        formats = [kind.upper() for kind in style.formats]
        vector = [kind for kind in formats if kind in ("PDF", "SVG", "EPS")]
        raster = [kind for kind in formats if kind not in vector]
        saved = []
        if vector:
            saved.append(f"as {_join(vector)} with editable text")
        if raster:
            saved.append(f"as {_join(raster)} at {style.dpi} dpi")
        drawn += (
            f", {style.width_mm:g} mm wide with text in {style.fonts[0]} at {style.font_size:g} pt, and "
            f"saved {' and '.join(saved)}"
        )
    sentences = [drawn + "."]
    if any(path.name.endswith(".source.csv") for asset in figures for path in asset.files):
        sentences.append("The data plotted in each figure were saved beside it as a CSV file (.source.csv).")
    sentences.append("What each figure shows, from its caption:")
    draft.paragraph(*sentences)
    draft.bullets(
        _figure_caption(asset) + (f" {references.cite(*asset.references)}" if asset.references else "")
        for asset in figures
    )


def _figure_caption(asset: Asset) -> str:
    head = f"{asset.number}. {asset.title}" if asset.number else asset.title
    body = asset.caption[len(head) :].strip() if asset.caption.startswith(head) else asset.caption
    body = body.removeprefix(".").strip()
    return f"**{asset.number or asset.id}.** {_sentence(asset.title)}" + (f" {body}" if body else "")


def _software_section(draft: Draft, result, references: References) -> None:
    from .. import __version__  # noqa: PLC0415

    study = result.study
    versions = provenance.package_versions()
    draft.heading("Software")
    python = f"Python {platform.python_version()} {references.cite(PYTHON_REFERENCE)}"
    ada = f"ADA {__version__} {references.cite(ADA_REFERENCE.format(version=__version__))}"
    used = list(LIBRARIES)
    if Path(study.data_path).suffix.lower() in EXCEL_SUFFIXES:
        used.append(("openpyxl", "openpyxl", ""))
    named = []
    for distribution, name, reference in used:
        if distribution in versions:
            cite = references.cite(reference) if reference else ""
            named.append(f"{name} {versions[distribution]}" + (f" {cite}" if cite else ""))
    draft.paragraph(
        f"The analysis was run in {python} with {ada}" + (f", which used {_join(named)}." if named else "."),
        "The file provenance.json records these versions and the checksum of every input and output file. "
        "Every table and figure can be rebuilt from the original data file and study.yml with `ada build`, "
        "and `ada verify` checks that a rebuild gives the same numbers.",
    )


# The results draft.


def results_draft(result) -> Draft:
    """One plain sentence per analysis with its estimates, intervals and p-values; exploratory ones apart."""
    study = result.study
    draft = Draft("Results (draft)")
    draft.note(
        "Draft for the authors to edit: one sentence per analysis, with the numbers this build computed. "
        "Planned analyses come first; exploratory analyses follow in their own paragraph. Put the sentences "
        "in the order the paper needs and give each number its unit. `ada build` rewrites this file."
    )
    entries = _analyses(result)
    names = _names(entries)
    done = [(spec, found) for spec, found in entries if found is not None and found.supported]
    planned = [_result_sentence(found, study, names) for spec, found in done if spec.planned]
    exploratory = [_result_sentence(found, study, names) for spec, found in done if not spec.planned]
    if not done:
        draft.paragraph("No analyses were run, so there are no results to report.")
    if planned:
        draft.paragraph(*planned)
    if exploratory:
        draft.paragraph(
            "The following analyses were exploratory: they were not planned before the data were seen, and "
            "their results suggest hypotheses rather than test them.",
            *exploratory,
        )
    skipped = [(spec, found) for spec, found in entries if found is None or not found.supported]
    if skipped:
        draft.paragraph(
            *[
                f"{_cap(_phrase(spec, names))} was not analysed."
                + (f" {_sentence(found.reason)}" if found is not None and found.reason else "")
                for spec, found in skipped
            ]
        )
    return draft


def _result_sentence(found, study, names: dict[str, str]) -> str:
    conf = _percent(study.decisions.confidence)
    size = f"n = {found.n:,}"
    if found.subjects:
        size += f" rows from {found.subjects:,} subjects"
    elif found.clusters:
        size += f" rows in {found.clusters:,} clusters"
    context = f"{_test_name(found.test)}; {size}" if found.test else size
    method = study.decisions.multiple_comparisons
    adjusted = ADJUSTMENTS[method][1] if method in ADJUSTMENTS else ""
    estimates = found.estimates if found.estimates is not None else pd.DataFrame()
    parts = []
    for row in estimates.to_dict("records"):
        label = row.get("label") or row.get("term") or "the estimate"
        parts.append((str(label), _estimate_text(row, found.statistic_name, conf, adjusted)))
    head = f"For {_phrase(found.spec, names)} ({context})"
    if not parts:
        sentence = f"{head}, the analysis gave no estimates."
    elif len(parts) == 1:
        sentence = f"{head}, the estimate for {parts[0][0]} was {parts[0][1]}."
    else:
        listed = _join([f"{label}, {text}" for label, text in parts], separator="; ")
        sentence = f"{head}, the estimates were: {listed}."
    effect = found.effect or {}
    if effect.get("name") and _finite(effect.get("value")):
        interval = (
            f" ({conf} CI {number(effect['ci_low'], 2)} to {number(effect['ci_high'], 2)})"
            if _finite(effect.get("ci_low")) and _finite(effect.get("ci_high"))
            else ""
        )
        sentence += f" {effect['name']} was {number(effect['value'], 2)}{interval}."
    return sentence


def _estimate_text(row: dict, statistic_name: str, conf: str, adjusted: str) -> str:
    estimate, low, high = row.get("estimate"), row.get("ci_low"), row.get("ci_high")
    decimals = _decimals(estimate, low, high)
    text = number(estimate, decimals) if _finite(estimate) else "not estimable"
    details = []
    if _finite(low) and _finite(high):
        details.append(f"{conf} CI {number(low, decimals)} to {number(high, decimals)}")
    test = []
    if _finite(row.get("statistic")):
        test.append(f"{statistic_name or 'statistic'} = {number(row['statistic'], 2)}")
    if _finite(row.get("df")):
        df = float(row["df"])
        test.append(f"df = {int(df)}" if df.is_integer() else f"df = {number(df, 1)}")
    if _finite(row.get("p")):
        test.append(_p(row["p"]))
    if test:
        details.append(", ".join(test))
    if adjusted and _finite(row.get("p_adjusted")):
        details.append(f"{adjusted}-adjusted {_p(row['p_adjusted'])}")
    return text + (f" ({'; '.join(details)})" if details else "")


def _decimals(estimate, low, high) -> int:
    """Decimals that show an estimate to the precision of its interval: about two significant figures."""
    if _finite(low) and _finite(high) and high > low:
        scale = (high - low) / 2
    elif _finite(estimate) and estimate:
        scale = abs(float(estimate))
    else:
        return 2
    return int(min(4, max(0, 1 - math.floor(math.log10(scale)))))


def _p(value) -> str:
    text = p_value(value)
    return f"p < {text[1:]}" if text.startswith("<") else f"p = {text}"


# The supplement.


def supplement_draft(result) -> Draft:
    """The supplementary tables and figures, the audit's items to check, the processing record, model
    output and the data dictionary, in one document."""
    study = result.study
    draft = Draft(f"Supplementary material: {study.title}")
    draft.paragraph(
        f"Built by ADA from {study.data_path} and study.yml. Tables and figures are numbered as in the "
        "paper's supplement. `ada build` rewrites this file, so edit a copy."
    )

    draft.heading("Supplementary tables")
    supplementary = [table for table in result.tables if table.number.startswith("Table S")]
    for table in supplementary:
        draft.add_table(table)
    if not supplementary:
        draft.paragraph("There are no supplementary tables.")

    draft.heading("Supplementary figures")
    figures = _figures(result, "supplement")
    for asset in figures:
        png = next((path for path in asset.files if path.suffix.lower() == ".png"), None)
        if png is not None and png.is_file():
            draft.picture(png, asset.title)
        draft.paragraph(_figure_caption(asset))
        draft.paragraph("Files: " + ", ".join(f"figures/{path.name}" for path in asset.files) + ".")
    if not figures:
        draft.paragraph("There are no supplementary figures.")

    draft.heading("Data audit")
    findings = getattr(result, "findings", []) or []
    checks = [finding for finding in findings if finding.status == "check"]
    if checks:
        draft.paragraph(
            f"The audit marked {_count(len(checks), 'item')} for checking. It changes nothing: an item that "
            "should change the data needs a processing step with a reason."
        )
        draft.bullets(
            (f"**{finding.column}**: " if finding.column else "")
            + finding.finding
            + (f" ({finding.where})" if finding.where else "")
            for finding in checks
        )
    else:
        draft.paragraph("The audit marked nothing for checking.")
    notes = sum(1 for finding in findings if finding.status == "note")
    if notes:
        draft.paragraph(f"Its {_count(notes, 'note')} for the record are in audit/data-audit.md.")

    draft.heading("Processing record")
    draft.add_table(processing_table(result))

    draft.heading("Model output")
    _model_output(draft, result)

    draft.heading("Data dictionary")
    dictionary = next((table for table in result.tables if table.id == "data-dictionary"), None)
    if dictionary is not None:
        draft.add_table(dictionary)
    else:
        draft.paragraph("No data dictionary was built.")
    return draft


def processing_table(result) -> Table:
    """The processing record as a table: each step with its reason and the rows before and after."""
    processed = result.processed
    columns = ["Step", "What", "Reason", "Rows before", "Rows after", "Values changed"]
    data = pd.DataFrame(
        [
            [
                record.number,
                record.description,
                record.reason,
                record.rows_before,
                record.rows_after,
                record.values_changed,
            ]
            for record in processed.records
        ],
        columns=columns,
    )
    display = data.astype(object)
    for column in ("Rows before", "Rows after", "Values changed"):
        display[column] = data[column].map(lambda value: f"{value:,}")
    excluded = len(processed.excluded) if processed.excluded is not None else 0
    notes = [
        f"The original file, {result.study.data_path} (SHA-256 beginning {processed.sha256[:12]}), was not "
        f"changed. {len(processed.data):,} rows remained for analysis."
    ]
    if excluded:
        notes.append(
            f"The {excluded:,} excluded rows are in processing/excluded-rows.csv, each with the step that "
            "removed it and its line in the original file."
        )
    return Table(
        id="processing-record",
        title="Steps from the original file to the analysed data",
        display=display.astype(str),
        data=data,
        notes=notes,
    )


def _model_output(draft: Draft, result) -> None:
    folder = result.out / "supplement" / "models"
    files = {path.stem: path for path in sorted(folder.glob("*.txt"))} if folder.is_dir() else {}
    entries = _analyses(result)
    if not entries and not files:
        draft.paragraph("No analyses were run, so there is no model output.")
        return
    for spec, found in entries:
        draft.heading(f"{_title(spec)} ({spec.id})", 3)
        status = "Planned" if spec.planned else "Exploratory"
        if found is None:
            draft.paragraph(f"{status}. The build produced no result for this analysis.")
            continue
        if not found.supported:
            draft.paragraph(f"{status}. Not analysed.", _sentence(found.reason) if found.reason else "")
            continue
        size = f"n = {found.n:,}" + (
            f" ({found.n_missing:,} left out for missing values)" if found.n_missing else ""
        )
        method = f" (`{found.method}`)" if found.method else ""
        draft.paragraph(f"{status}. {found.test or 'Test not named'}{method}; {size}.")
        if found.formula:
            draft.paragraph(f"Model: `{found.formula}`")
        if found.diagnostics:
            draft.paragraph("Assumption checks:")
            draft.bullets(found.diagnostics)
        if found.warnings:
            draft.paragraph("Caveats:")
            draft.bullets(found.warnings)
        path = files.pop(spec.id, None)
        text = path.read_text(encoding="utf-8") if path is not None else found.summary
        if text and text.strip():
            if path is not None:
                draft.paragraph(f"Output (supplement/models/{path.name}):")
            draft.code(text)
    for stem, path in files.items():
        draft.heading(stem, 3)
        draft.paragraph(f"Output (supplement/models/{path.name}):")
        draft.code(path.read_text(encoding="utf-8"))


# Writing a draft.


def write(draft: Draft, stem: Path, folder: Path) -> list[Path]:
    """Write ``draft`` as ``stem``.md and ``stem``.docx; pictures are linked relative to ``folder``."""
    markdown_path = stem.with_suffix(".md")
    markdown_path.write_text(to_markdown(draft, folder), encoding="utf-8")
    docx_path = stem.with_suffix(".docx")
    to_docx(draft, docx_path)
    return [markdown_path, docx_path]


def to_markdown(draft: Draft, folder: Path) -> str:
    parts = [f"# {draft.title}"]
    for block in draft.blocks:
        if block.kind == "heading":
            parts.append(f"{'#' * block.level} {block.text}")
        elif block.kind == "paragraph":
            parts.append(block.text)
        elif block.kind == "note":
            parts.append(f"> {block.text}")
        elif block.kind == "bullets":
            parts.append("\n".join(f"- {item}" for item in block.items))
        elif block.kind == "numbered":
            parts.append("\n".join(f"{index}. {item}" for index, item in enumerate(block.items, 1)))
        elif block.kind == "table":
            parts.append(markdown(block.table).rstrip("\n"))
        elif block.kind == "code":
            parts.append(f"```\n{block.text}\n```")
        elif block.kind == "image":
            relative = Path(os.path.relpath(Path(block.image).resolve(), folder.resolve())).as_posix()
            parts.append(f"![{block.text}]({relative})")
    return "\n\n".join(parts) + "\n"


def to_docx(draft: Draft, path: Path) -> None:
    """Write ``draft`` as a Word file with real headings, lists, tables and pictures."""
    from docx.shared import Mm, Pt  # noqa: PLC0415

    document = docx_document()
    document.add_heading(draft.title, level=0)
    for block in draft.blocks:
        if block.kind == "heading":
            document.add_heading(block.text, level=max(1, block.level - 1))
        elif block.kind == "paragraph":
            _runs(document.add_paragraph(), block.text)
        elif block.kind == "note":
            paragraph = document.add_paragraph()
            _runs(paragraph, block.text)
            for run in paragraph.runs:
                run.italic = True
        elif block.kind == "bullets":
            for item in block.items:
                _runs(document.add_paragraph(style="List Bullet"), item)
        elif block.kind == "numbered":
            for index, item in enumerate(block.items, 1):
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Mm(7)
                paragraph.paragraph_format.first_line_indent = Mm(-7)
                _runs(paragraph, f"{index}.\t{item}")
        elif block.kind == "table":
            add_docx_table(document, block.table)
            document.add_paragraph()
        elif block.kind == "code":
            paragraph = document.add_paragraph()
            lines = block.text.split("\n")
            for index, line in enumerate(lines):
                run = paragraph.add_run(line)
                run.font.name = "Courier New"
                run.font.size = Pt(7)
                if index < len(lines) - 1:
                    run.add_break()
        elif block.kind == "image":
            picture = document.add_picture(str(block.image))
            if picture.width > Mm(160):
                picture.height = int(picture.height * Mm(160) / picture.width)
                picture.width = Mm(160)
    save_docx(document, path)


def _runs(paragraph, text: str) -> None:
    """Add ``text`` to a Word paragraph, turning **bold** and `code` into formatted runs."""
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            paragraph.add_run(part[1:-1]).font.name = "Courier New"
        else:
            paragraph.add_run(part)


# Small helpers.


def _analyses(result) -> list[tuple]:
    """Each analysis in study.yml order with its result (None if the build made none), then any others."""
    results = getattr(result, "results", {}) or {}
    entries = [(spec, results.get(spec.id)) for spec in result.study.analyses]
    listed = {spec.id for spec in result.study.analyses}
    entries += [(found.spec, found) for ident, found in sorted(results.items()) if ident not in listed]
    return entries


def _names(entries) -> dict[str, str]:
    """Each analysis's title, with its id added where two analyses would otherwise share a title."""
    titles = [_title(spec) for spec, _ in entries]
    return {
        spec.id: title + (f" (`{spec.id}`)" if titles.count(title) > 1 else "")
        for (spec, _), title in zip(entries, titles, strict=True)
    }


def _record(processed, kind: str):
    return next((record for record in processed.records if record.kind == kind), None)


def _asset(result, ident: str) -> Asset | None:
    return next((asset for asset in result.assets if asset.id == ident), None)


def _figures(result, section: str) -> list[Asset]:
    return [asset for asset in result.assets if asset.kind == "figure" and asset.section == section]


def _title(spec) -> str:
    if spec.kind == "compare":
        return f"{spec.outcome} by {spec.compare}"
    if spec.kind == "relate":
        return f"Association between {spec.outcome} and {spec.relate}"
    return f"Change in {spec.outcome} over {spec.change}" + (f" by {spec.by}" if spec.by else "")


def _phrase(spec, names: dict[str, str]) -> str:
    """An analysis's name as it reads inside a sentence: "the change in Weight over Day"."""
    name = names.get(spec.id, _title(spec))
    return f"the {name[0].lower()}{name[1:]}" if spec.kind != "compare" else name


def _test_name(test: str) -> str:
    """A test's name inside a sentence: "linear mixed model", but "Welch's t-test" and "ANOVA" as they are."""
    text = test.strip()
    first = text.split()[0] if text else ""
    if not first or "'" in first or "’" in first or first.lower().split("-")[0] in _NAMED_TESTS:
        return text
    if len(first) > 1 and first[:2].isupper():
        return text
    return text[0].lower() + text[1:]


def _with_article(test: str) -> str:
    """A test's name as it reads after "with": "Welch's t-test", "a linear mixed model"."""
    text = _test_name(test)
    if not text:
        return "a method the build did not name"
    first = text.split()[0]
    if "'" in first or "’" in first:
        return text
    if first.lower().split("-")[0] in _NAMED_TESTS:
        return f"the {text}"
    if len(first) > 1 and first[:2].isupper():
        return ("an " if first[0] in "AEFHILMNORSX" else "a ") + text
    vowel = text[0] in "aeiou" and not text.startswith(("one", "uni", "use", "eu"))
    return ("an " if vowel else "a ") + text


def _count(count: int, noun: str, plural: str | None = None) -> str:
    """A count of things in the text, in words below ten as journals print them: "two analyses"."""
    count = int(count)
    word = _WORDS[count] if 0 <= count < len(_WORDS) else f"{count:,}"
    return f"{word} {noun if count == 1 else plural or noun + 's'}"


def _n(count: int, noun: str) -> str:
    """A count of data, in figures: "5 rows", "19 subjects", or "no rows"."""
    count = int(count)
    return f"{'no' if count == 0 else f'{count:,}'} {noun if count == 1 else noun + 's'}"


def _was(count: int) -> str:
    return "was" if count == 1 else "were"


def _cap(text: str) -> str:
    return text[:1].upper() + text[1:]


def _join(items, word: str = "and", separator: str = ", ") -> str:
    """ "A, B and C"; with another separator, "A; B; and C" so each item can hold commas."""
    items = [str(item) for item in items]
    if len(items) <= 1:
        return "".join(items)
    tail = f" {word} " if separator == ", " else f"{separator.rstrip()} {word} "
    return separator.join(items[:-1]) + tail + items[-1]


def _sentence(text: str) -> str:
    text = str(text).strip()
    return text if not text or text[-1] in ".!?" else text + "."


def _spread(counts: pd.Series) -> str:
    low, high = int(counts.min()), int(counts.max())
    return f"{low}" if low == high else f"{low} to {high}"


def _moments(data: pd.DataFrame, time: str) -> list:
    if isinstance(data[time].dtype, pd.CategoricalDtype):
        return levels_of(data, time)
    return sorted(data[time].dropna().unique())


def _value(value) -> str:
    if isinstance(value, pd.Timestamp):
        return f"{value:%Y-%m-%d}"
    if isinstance(value, float | int) or hasattr(value, "item"):
        number_ = float(value)
        return f"{int(number_)}" if number_.is_integer() else f"{number_:g}"
    return str(value)


def _percent(fraction: float) -> str:
    return f"{fraction * 100:g}%"


def _finite(value) -> bool:
    try:
        return value is not None and math.isfinite(float(value))
    except (TypeError, ValueError):
        return False
