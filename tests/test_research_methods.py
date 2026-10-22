import contextlib
import json
import platform
import unittest
from dataclasses import dataclass
from unittest import mock

import pandas as pd
from docx import Document
from research_helpers import assay, assay_study, study_folder, weights

from ada import __version__
from ada.research import build, provenance
from ada.research.analyses import COLUMNS, REFERENCES
from ada.research.methods import LIBRARIES, _for_paper, _join, _test_prose, _with_article
from ada.research.results import ESTIMATE_COLUMNS, AnalysisResult
from ada.research.tables import Table

WELCH = "Welch, B. L. (1947). The generalization of 'Student's' problem. Biometrika 34, 28–35."
STEPS = [
    {
        "exclude": {"where": {"column": "Viability (%)", "below": 0}},
        "reason": "Readings below zero are errors.",
    },
    {"exclude": {"ids": ["C00"], "id_column": "Well"}, "reason": "Contaminated well."},
    {"transform": {"column": "Passage", "to": "log_passage", "apply": "log10"}, "reason": "Skewed."},
]


def assay_with_steps(**changes) -> dict:
    study = assay_study(processing=STEPS, **changes)
    study["variables"]["log_passage"] = {"role": "other", "type": "numeric"}
    return study


def estimates(*rows) -> pd.DataFrame:
    """Rows in the statistics step's shape: ESTIMATE_COLUMNS plus key, measure, statistic_name, df2, test
    and adjustment."""
    return pd.DataFrame(
        list(rows), columns=COLUMNS if any("measure" in row for row in rows) else ESTIMATE_COLUMNS
    )


def welch(spec, result) -> AnalysisResult:
    """Welch's ANOVA with Games–Howell contrasts, shaped as the statistics step writes it; the counts come
    from the processed data."""
    used = result.processed.rows_used_by([spec.outcome, spec.compare])
    frame = result.data[used]
    return AnalysisResult(
        spec=spec,
        design=result.study.design,
        test="Welch's ANOVA with Games–Howell contrasts",
        method="statsmodels.stats.oneway.anova_oneway(use_var='unequal')",
        n=int(used.sum()),
        n_missing=int((~used).sum()),
        groups={
            str(level): int((frame[spec.compare] == level).sum()) for level in ("Control", "Low", "High")
        },
        estimates=estimates(
            {
                "term": "Treatment",
                "label": "Treatment (overall)",
                "statistic": 24.31,
                "df": 2.0,
                "df2": 21.7,
                "p": 3e-6,
                "p_adjusted": 3e-6,
                "key": True,
                "measure": "overall",
                "statistic_name": "F",
                "test": "Welch's ANOVA",
                "adjustment": "Holm over family primary (1 analysis)",
            },
            {
                "term": "High",
                "label": "High − Control",
                "estimate": -12.08,
                "ci_low": -15.18,
                "ci_high": -8.98,
                "statistic": -8.12,
                "df": 21.46,
                "p": 2e-7,
                "p_adjusted": 6e-7,
                "key": False,
                "measure": "mean difference",
                "statistic_name": "t",
                "test": "Games–Howell",
                "adjustment": "Games–Howell over the 3 pairs",
            },
        ),
        statistic_name="F",
        effect={"name": "omega squared", "value": 0.61, "ci_low": 0.3, "ci_high": 0.75},
        diagnostics=[
            "1 of 34 rows were left out because a value this analysis needs was missing (missing values: "
            "Viability (%) 1).",
            "Rows per group: Control (n = 10), Low (n = 11), High (n = 12).",
            "Shapiro–Wilk test of the residuals: W = 0.970, p = 0.410 (n = 33), consistent with a normal "
            "distribution.",
            "Levene's test (median-centred) for equal variances across Treatment: W = 1.20, p = 0.310.",
        ],
        references=(WELCH, REFERENCES["shapiro"], REFERENCES["omega"], REFERENCES["holm"]),
    )


def regression(spec, result) -> AnalysisResult:
    return AnalysisResult(
        spec=spec,
        design=result.study.design,
        test="Linear regression (OLS) adjusted for Treatment",
        method="statsmodels.formula.api.ols",
        formula="Viability (%) ~ Passage + Treatment",
        n=33,
        n_missing=1,
        estimates=estimates(
            {
                "term": "Passage",
                "label": "Slope on Passage",
                "estimate": 0.21,
                "ci_low": -0.35,
                "ci_high": 0.77,
                "statistic": 0.76,
                "df": 29.0,
                "p": 0.45,
                "p_adjusted": 0.45,
                "key": True,
                "measure": "slope",
                "statistic_name": "t",
                "test": "Linear regression (OLS)",
                "adjustment": "Holm over family exploratory (1 analysis)",
            },
            {
                "term": "Intercept",
                "label": "Intercept",
                "estimate": 85.2,
                "ci_low": 80.1,
                "ci_high": 90.3,
                "statistic": 33.0,
                "df": 29.0,
                "p": 1e-20,
                "key": False,
                "measure": "intercept",
                "statistic_name": "t",
                "test": "Linear regression (OLS)",
            },
        ),
        statistic_name="t",
        diagnostics=[
            "33 rows with both Viability (%) and Passage.",
            "Breusch–Pagan test for constant residual variance: LM = 1.20, p = 0.550.",
            "The design's time variable (Day) is not in this model, so the difference is averaged over time "
            "points; a `change: Day` analysis models time.",
        ],
        warnings=["Passage was recorded for each well, not each plate."],
        references=(REFERENCES["breusch_pagan"], REFERENCES["holm"]),
        summary="OLS Regression Results\nR-squared: 0.81",
    )


def unsupported(spec, result) -> AnalysisResult:
    return AnalysisResult(
        spec=spec,
        design=result.study.design,
        supported=False,
        reason="Binary outcomes in repeated designs are not fitted; write your own analysis under `custom:`.",
    )


def analyses_made_by(**makers):
    """Stand in for the statistics step: each analysis id gets the result its maker returns."""

    def run(result):
        for spec in result.study.analyses:
            if spec.id in makers:
                result.results[spec.id] = makers[spec.id](spec, result)

    return mock.patch("ada.research.analyses.run", run)


def two_analyses() -> list[dict]:
    return [
        {
            "id": "primary",
            "outcome": "Viability (%)",
            "compare": "Treatment",
            "planned": True,
            "family": "primary",
        },
        {
            "id": "passage",
            "outcome": "Viability (%)",
            "relate": "Passage",
            "covariates": ["Treatment"],
            "planned": False,
            "family": "exploratory",
        },
    ]


@dataclass
class Made:
    """A build's drafts, read before its scratch folder goes."""

    built: object
    methods: str
    results: str
    supplement: str
    record: dict
    methods_docx: object
    supplement_docx: object


def make(study: dict, frame=None, data_name: str = "assay.csv", draw=None, **makers) -> Made:
    """Build ``study`` with stand-in analysis results (and figures) and read what the methods step wrote."""
    drawing = mock.patch("ada.research.plots.draw_all", draw) if draw else contextlib.nullcontext()
    data = assay() if frame is None else frame
    with study_folder(data, study, data_name) as path, analyses_made_by(**makers), drawing:
        built = build(path)
        return Made(
            built=built,
            methods=read(built, "methods/methods.md"),
            results=read(built, "methods/results.md"),
            supplement=read(built, "supplement/supplement.md"),
            record=json.loads(read(built, "provenance.json")),
            methods_docx=Document(str(built.out / "methods" / "methods.docx")),
            supplement_docx=Document(str(built.out / "supplement" / "supplement.docx")),
        )


def read(built, name: str) -> str:
    return (built.out / name).read_text(encoding="utf-8")


class MethodsDraftTests(unittest.TestCase):
    def test_each_processing_step_is_named_with_its_counts_and_reason(self):
        text = make(assay_with_steps(), primary=welch).methods
        self.assertIn("The data were read from assay.csv, which held 36 rows and four columns", text)
        self.assertIn("**Step 1.** assay.csv was read: 36 rows read and 36 kept", text)
        self.assertIn(
            "**Step 2.** The missing-value codes listed in study.yml (-99 in every column) were read as "
            "missing",
            text,
        )
        self.assertIn(
            "**Step 3.** Rows in which Viability (%) was below 0 were excluded (1 row removed, leaving 35). "
            "Reason: Readings below zero are errors.",
            text,
        )
        self.assertIn(
            "**Step 4.** Rows in which Well was C00 were excluded (1 row removed, leaving 34). "
            "Reason: Contaminated well.",
            text,
        )
        self.assertIn(
            "**Step 5.** log_passage was set to the base-10 logarithm of Passage (34 values computed)", text
        )
        self.assertIn("After processing, 34 of the 36 rows read remained for analysis.", text)
        self.assertIn("The original file was not changed, and the 2 excluded rows were kept", text)
        self.assertIn("Missing values were coded as -99 in every column.", text)

    def test_each_analysis_is_named_with_its_test_status_and_missing_rows(self):
        study = assay_with_steps(analyses=two_analyses())
        made = make(study, primary=welch, passage=regression)
        text, primary = made.methods, made.built.results["primary"]
        self.assertIn("listed two analyses, and no other analyses were run", text)
        self.assertIn(
            "Viability (%) was compared between the levels of Treatment (Control, 10 rows; Low, 11 rows; "
            "High, 12 rows) with Welch's ANOVA with Games–Howell contrasts [",
            text,
        )
        self.assertIn(
            f"The analysis used the {primary.n} rows that had every value it needed (complete cases); "
            f"{primary.n_missing} row was left out because a value was missing.",
            text,
        )
        self.assertIn("This analysis was planned before the data were seen.", text)
        self.assertIn("This analysis was exploratory: it was not planned before the data were seen.", text)
        self.assertIn("The model was `Viability (%) ~ Passage + Treatment`.", text)
        self.assertIn("Treatment entered the model as a fixed-effect covariate.", text)
        self.assertIn("Omega squared was reported as the standardised effect size", text)
        self.assertIn(
            "Its key result, the one adjusted across its family, was the overall test of Treatment.", text
        )
        self.assertIn(
            "The p-values of its other rows were adjusted within the analysis (Games–Howell over the 3 "
            "pairs), with simultaneous confidence intervals.",
            text,
        )
        checks = text.split("Its fit and assumptions were checked [", 1)[1].split("\n", 1)[0]
        self.assertIn("Shapiro–Wilk test of the residuals: W = 0.970", checks)
        self.assertIn("Levene's test (median-centred)", checks)
        self.assertNotIn("Rows per group", text)
        self.assertNotIn("this analysis needs was missing", text)
        self.assertNotIn("rows with both", text)
        self.assertIn("Breusch–Pagan test for constant residual variance", text)
        self.assertIn("so the difference is averaged over time points.", text)
        self.assertNotIn("`change: Day`", text)
        self.assertIn("Caveat: Passage was recorded for each well, not each plate.", text)
        self.assertIn(
            "each analysis contributed the p-value of its key result, adjusted with Holm's step-down", text
        )
        self.assertIn("All tests were two-sided.", text)
        self.assertIn("“primary” (Viability (%) by Treatment) and “exploratory” (Association between", text)
        self.assertIn(
            "The significance level was α = 0.05, and all confidence intervals were computed at the 95% "
            "level.",
            text,
        )
        self.assertIn("Exploratory results are labelled as such", text)
        self.assertIn(WELCH, text)
        references = text.split("## References", 1)[1]
        self.assertEqual(references.count("Holm, S. (1979)"), 1)

    def test_the_software_section_names_python_and_every_package_version(self):
        text = make(assay_study(), primary=welch).methods
        versions = provenance.package_versions()
        self.assertIn(f"Python {platform.python_version()} [", text)
        self.assertIn(f"ADA {__version__} [", text)
        for distribution, name, _ in LIBRARIES:
            if distribution in versions:
                self.assertIn(f"{name} {versions[distribution]}", text)
        references = text.split("## References", 1)[1]
        self.assertIn("1. ", references)
        self.assertIn("Harris, C. R.", references)
        self.assertIn("Holm, S. (1979)", references)

    def test_the_settings_in_study_yml_change_the_adjustment_and_levels(self):
        decisions = {"multiple_comparisons": "bonferroni", "alpha": 0.01, "confidence": 0.99}
        made = make(assay_study(decisions=decisions), primary=welch)
        text = made.methods
        self.assertIn("with the Bonferroni correction [", text)
        self.assertIn("Dunn, O. J. (1961)", text)
        self.assertIn("α = 0.01, and all confidence intervals were computed at the 99% level", text)
        self.assertIn("(99% CI −15.2 to −9.0;", made.results)

    def test_unsupported_analyses_are_explained(self):
        made = make(assay_study(), primary=unsupported)
        methods, results = made.methods, made.results
        self.assertIn(
            "This planned analysis was not run. Binary outcomes in repeated designs are not fitted.", methods
        )
        self.assertNotIn("custom:", methods)
        self.assertNotIn("Holm", methods)
        self.assertIn("Viability (%) by Treatment was not analysed. Binary outcomes", results)
        self.assertIn("No analyses were run, so there are no results to report.", results)
        self.assertIn("Planned. Not analysed. Binary outcomes", made.supplement)

    def test_a_study_without_analyses_says_so(self):
        made = make(assay_study(analyses=[]))
        text = made.methods
        self.assertIn("No statistical analyses were listed in study.yml and none were run", text)
        self.assertNotIn("Figure 1.", text)
        self.assertIn("No analyses were run, so there is no model output.", made.supplement)

    def test_an_analysis_the_build_did_not_fit_is_not_described_as_run(self):
        text = make(assay_study()).methods
        self.assertIn(
            "This planned analysis was listed in study.yml, but the build produced no result for it.", text
        )
        self.assertNotIn("compared between", text)

    def test_a_repeated_design_is_described_with_its_counts(self):
        study = {
            "study": {"title": "Weights"},
            "data": {"path": "weights.csv", "observation": "one mouse on one day"},
            "design": {"type": "repeated", "subject": "Mouse", "time": "Day"},
            "variables": {
                "Mouse": {"role": "identifier", "type": "text"},
                "Diet": {"role": "group", "type": "categorical", "levels": ["Chow", "High fat"]},
                "Day": {"role": "time", "type": "numeric", "unit": "days"},
                "Weight (g)": {"role": "outcome", "type": "numeric", "unit": "g"},
            },
            "analyses": [
                {"id": "growth", "outcome": "Weight (g)", "change": "Day", "by": "Diet", "planned": True},
                {
                    "id": "heavy",
                    "outcome": "Weight (g)",
                    "compare": "Diet",
                    "planned": False,
                    "test": "logistic",
                },
            ],
        }

        def mixed(spec, result):
            return AnalysisResult(
                spec=spec,
                design="repeated",
                test="Linear mixed model",
                method="statsmodels.formula.api.mixedlm",
                formula="Weight (g) ~ Day * Diet + (1 | Mouse)",
                n=60,
                subjects=12,
                groups={"Chow": 30, "High fat": 30},
                estimates=estimates(
                    {
                        "term": "Day:Diet[T.High fat]",
                        "label": "Day × High fat",
                        "estimate": 0.0712,
                        "ci_low": 0.0601,
                        "ci_high": 0.0823,
                        "statistic": 12.6,
                        "p": 1e-12,
                        "p_adjusted": 1e-12,
                        "key": True,
                        "measure": "slope difference",
                        "statistic_name": "z",
                        "test": "Linear mixed model",
                        "adjustment": "Holm over family growth (1 analysis)",
                    }
                ),
                statistic_name="z",
            )

        made = make(study, weights(), "weights.csv", growth=mixed, heavy=unsupported)
        text = made.methods
        self.assertIn(
            "The design was repeated measures: the 60 rows analysed came from 12 subjects, identified by "
            "Mouse, with 5 rows each.",
            text,
        )
        self.assertIn(
            "Measurements were made at five time points (Day 0, 7, 14, 21 and 28), with 12 subjects", text
        )
        self.assertIn("By Diet, 6 subjects were in Chow and 6 in High fat.", text)
        self.assertIn(
            "The change in Weight (g) over Day, and whether it differed between the levels of Diet (Chow, 30 "
            "rows; High fat, 30 rows), was estimated with a linear mixed model",
            text,
        )
        self.assertIn(
            "The analysis used 60 rows, from 12 subjects; none was left out for a missing value.", text
        )
        self.assertIn(
            "For the change in Weight (g) over Day by Diet (linear mixed model; n = 60 rows from 12 "
            "subjects), the estimate for Day × High fat was 0.0712 g per day (95% CI 0.0601 to 0.0823; "
            "z = 12.60, p < 0.001; adjusted p < 0.001, Holm over family growth, 1 analysis).",
            made.results,
        )

    def test_two_builds_give_identical_markdown(self):
        study = assay_with_steps(analyses=two_analyses())
        with study_folder(assay(), study) as path, analyses_made_by(primary=welch, passage=regression):
            first = build(path)
            texts = {
                name: read(first, name)
                for name in ("methods/methods.md", "methods/results.md", "supplement/supplement.md")
            }
            second = build(path)
            for name, text in texts.items():
                self.assertEqual(read(second, name), text, name)

    def test_the_word_files_open_and_hold_the_headings(self):
        made = make(assay_with_steps(analyses=two_analyses()), primary=welch, passage=regression)
        methods, supplement = made.methods_docx, made.supplement_docx
        headings = [p.text for p in methods.paragraphs if p.style.name.startswith(("Heading", "Title"))]
        self.assertEqual(
            headings,
            [
                "Methods (draft)",
                "Data",
                "Data processing",
                "Statistical analysis",
                "Figures",
                "Software",
                "References",
            ],
        )
        headings = [p.text for p in supplement.paragraphs if p.style.name.startswith(("Heading", "Title"))]
        for heading in (
            "Supplementary tables",
            "Data audit",
            "Processing record",
            "Model output",
            "Data dictionary",
        ):
            self.assertIn(heading, headings)
        self.assertEqual(supplement.core_properties.author, "")
        self.assertEqual(supplement.core_properties.created.year, 2000)
        body = "\n".join(p.text for p in methods.paragraphs)
        self.assertIn("Welch's ANOVA", body)
        self.assertNotIn("**", body)


class ResultsDraftTests(unittest.TestCase):
    def test_each_result_has_its_numbers_and_exploratory_ones_stand_apart(self):
        text = make(assay_with_steps(analyses=two_analyses()), primary=welch, passage=regression).results
        planned, exploratory = text.split("The following analyses were exploratory", 1)
        self.assertIn(
            "For Viability (%) by Treatment (Welch's ANOVA; n = 33), the overall test of Treatment gave "
            "F(2, 21.7) = 24.31 (p < 0.001; adjusted p < 0.001, Holm over family primary, 1 analysis) and "
            "the estimate for High − Control was −12.08 percentage points (95% CI −15.18 to −8.98; "
            "t(21.5) = −8.12, p < 0.001; adjusted p < 0.001, Games–Howell over the 3 pairs). Omega squared "
            "was 0.61 (95% CI 0.30 to 0.75).",
            planned,
        )
        self.assertNotIn("Passage", planned.split("> Draft", 1)[1].split("\n\n", 1)[1])
        self.assertIn("were not planned before the data were seen", exploratory)
        self.assertIn(
            "the estimate for Slope on Passage was 0.210 percentage points per unit of Passage "
            "(95% CI −0.350 to 0.770; "
            "t(29) = 0.76, p = 0.450; adjusted p = 0.450, Holm over family exploratory, 1 analysis).",
            exploratory,
        )
        self.assertNotIn("Intercept", text)
        self.assertNotIn("not estimable", text)
        self.assertNotIn("significan", text.lower())
        self.assertNotIn("trend", text.lower())


class SupplementTests(unittest.TestCase):
    def test_the_supplement_collects_tables_figures_audit_record_models_and_dictionary(self):
        def with_extras(spec, result):
            found = regression(spec, result)
            frame = pd.DataFrame({"Term": ["Passage"], "Estimate": ["0.21"]})
            result.add_table(
                Table(
                    id="passage-model", title="Viability and passage (exploratory)", display=frame, data=frame
                ),
                section="supplement",
                depends=("analyses",),
                calculation="test",
                code="test",
            )
            folder = result.folder("supplement") / "models"
            folder.mkdir(parents=True, exist_ok=True)
            (folder / "passage.txt").write_text("Model output for passage\n", encoding="utf-8")
            return found

        def draw(result):
            import matplotlib.pyplot as plt

            from ada.research.figures import save

            figure, axes = plt.subplots(figsize=(2, 2))
            axes.plot([0, 1], [0, 1])
            save(
                result,
                figure,
                ident="residuals",
                title="Residuals against fitted values",
                caption="Points are wells.",
                source=None,
                depends=("analyses",),
                code="test",
                section="supplement",
            )

        study = assay_with_steps(analyses=two_analyses())
        study["variables"]["Viability (%)"]["range"] = [0, 80]
        made = make(study, draw=draw, primary=welch, passage=with_extras)
        text, document, methods, record = made.supplement, made.supplement_docx, made.methods, made.record
        self.assertIn("**Table S1. Viability and passage (exploratory)**", text)
        self.assertIn("![Residuals against fitted values](../figures/residuals.png)", text)
        self.assertIn("**Figure S1.** Residuals against fitted values. Points are wells.", text)
        self.assertIn("Files: figures/residuals.pdf, figures/residuals.svg, figures/residuals.png.", text)
        self.assertIn("The audit marked one item for checking.", text)
        self.assertIn("fall outside the allowed range 0 to 80", text)
        self.assertIn(
            "| 3 | Exclude rows where Viability (%) is below 0 | Readings below zero are errors. | 36 | 35 "
            "| 0 |",
            text,
        )
        self.assertIn(
            "The model's record, supplement/models/passage.txt:\n\n```\nModel output for passage\n```", text
        )
        self.assertNotIn("Warnings:", text)
        self.assertIn("Diagnostics:\n\n- 1 of 34 rows were left out", text)
        self.assertIn("**Data dictionary**", text)
        self.assertEqual(len(document.inline_shapes), 1)
        self.assertEqual(len(document.tables), 3)
        self.assertIn("- **Figure S1.** Residuals against fitted values. Points are wells.", methods)

        for ident, section in (
            ("methods-draft", "record"),
            ("results-draft", "record"),
            ("supplement", "supplement"),
        ):
            entry = record["assets"][ident]
            self.assertEqual(entry["kind"], "text")
            self.assertEqual(entry["section"], section)
            self.assertIn("executed", entry["calculation"])
            self.assertIn("analyses", entry["depends_on"])
            self.assertIn("processing", entry["depends_on"])
        self.assertEqual(
            sorted(record["assets"]["supplement"]["files"]),
            ["supplement/supplement.docx", "supplement/supplement.md"],
        )


class WordingTests(unittest.TestCase):
    def test_tests_read_naturally_after_with(self):
        self.assertEqual(_with_article("Welch's t-test"), "Welch's t-test")
        self.assertEqual(_with_article("Linear mixed model"), "a linear mixed model")
        self.assertEqual(
            _with_article("Ordinary least squares regression"), "an ordinary least squares regression"
        )
        self.assertEqual(_with_article("One-way ANOVA"), "a one-way ANOVA")
        self.assertEqual(_with_article("Mann-Whitney U test"), "the Mann-Whitney U test")
        self.assertEqual(_with_article("ANOVA"), "an ANOVA")
        self.assertEqual(
            _with_article("Kruskal–Wallis test with pairwise Mann–Whitney tests"),
            "the Kruskal–Wallis test with pairwise Mann–Whitney tests",
        )
        self.assertEqual(
            _with_article(_test_prose("Linear mixed model, random intercept per subject (Mouse)")),
            "a linear mixed model with a random intercept per subject (Mouse)",
        )

    def test_advice_about_study_yml_stays_out_of_the_paper(self):
        self.assertEqual(
            _for_paper("Not fitted here. Write your own analysis under `custom:`."), "Not fitted here."
        )
        self.assertEqual(
            _for_paper("Time is not in this model, so it is averaged; a `change: Day` analysis models time."),
            "Time is not in this model, so it is averaged.",
        )
        self.assertEqual(
            _for_paper("`by` applies to change analyses only."), "`by` applies to change analyses only."
        )

    def test_lists_join_without_a_serial_comma(self):
        self.assertEqual(_join(["a"]), "a")
        self.assertEqual(_join(["a", "b"]), "a and b")
        self.assertEqual(_join(["a", "b", "c"]), "a, b and c")
        self.assertEqual(_join(["a, 1", "b, 2"], separator="; "), "a, 1; and b, 2")


if __name__ == "__main__":
    unittest.main()
