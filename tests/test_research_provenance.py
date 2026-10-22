"""Byte-stable builds, and the change report a rebuild writes to CHANGES.md."""

import io
import json
import re
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd
import yaml
from research_helpers import assay, assay_study, study_folder

from ada.research import build
from ada.research.provenance import compare, general_causes, step_changes
from ada.research.tables import DOCX_TIMESTAMP, stable_docx

EXCLUSION = {"exclude": {"where": {"column": "Viability (%)", "below": 0}}, "reason": "Instrument error."}
IMAGES = {".png", ".pdf", ".tif", ".tiff", ".eps", ".jpg"}


def snapshot(out: Path) -> dict[str, bytes]:
    """Every file a build wrote, except images, by path."""
    return {
        path.relative_to(out).as_posix(): path.read_bytes()
        for path in sorted(out.rglob("*"))
        if path.is_file() and path.suffix not in IMAGES
    }


def edit(path: Path, change) -> None:
    study = yaml.safe_load(path.read_text(encoding="utf-8"))
    change(study)
    path.write_text(yaml.safe_dump(study, sort_keys=False, allow_unicode=True), encoding="utf-8")


class DeterminismTests(unittest.TestCase):
    def test_building_twice_gives_the_same_bytes(self):
        with study_folder(assay(), assay_study(processing=[EXCLUSION])) as path:
            first = snapshot(build(path).out)
            second = snapshot(build(path).out)
            elsewhere = Path(tempfile.mkdtemp())
            try:
                shutil.copyfile(path, elsewhere / "study.yml")
                shutil.copyfile(path.parent / "assay.csv", elsewhere / "assay.csv")
                third = snapshot(build(elsewhere / "study.yml").out)
            finally:
                shutil.rmtree(elsewhere)
        self.assertIn("tables/descriptives.docx", first)
        self.assertIn("CHANGES.md", first)
        self.assertEqual(sorted(first), sorted(second))
        for name in first:
            self.assertEqual(first[name], second[name], f"{name} changed between two builds")
            self.assertEqual(first[name], third[name], f"{name} differs between two folders")

    def test_word_files_have_fixed_dates_inside(self):
        with study_folder(assay(), assay_study()) as path:
            out = build(path).out
            with zipfile.ZipFile(out / "tables" / "descriptives.docx") as archive:
                stamps = {info.date_time for info in archive.infolist()}
        self.assertEqual(stamps, {DOCX_TIMESTAMP})

    def test_stable_docx_keeps_the_content(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "note.docx"
            document = Document()
            document.add_paragraph("Viability fell with dose.")
            document.save(path)
            stable_docx(path)
            stable = path.read_bytes()
            stable_docx(path)
            self.assertEqual(path.read_bytes(), stable)
            self.assertEqual(Document(io.BytesIO(stable)).paragraphs[0].text, "Viability fell with dose.")

    def test_csv_numbers_are_written_to_full_precision_the_same_way_each_time(self):
        rng = np.random.default_rng(1)
        frame = pd.DataFrame({"x": rng.normal(size=200) * 10.0 ** rng.integers(-8, 8, 200), "y": 0.1 + 0.2})
        first = frame.to_csv(index=False, lineterminator="\n")
        self.assertEqual(first, frame.to_csv(index=False, lineterminator="\n"))
        self.assertIn("0.30000000000000004", first)
        back = pd.read_csv(io.StringIO(first), float_precision="round_trip")
        self.assertTrue((back["x"].to_numpy() == frame["x"].to_numpy()).all())

    def test_the_record_holds_no_dates_or_commit(self):
        with study_folder(assay(), assay_study()) as path:
            out = build(path).out
            text = (out / "provenance.json").read_text(encoding="utf-8")
            changes = (out / "CHANGES.md").read_text(encoding="utf-8")
        record = json.loads(text)
        self.assertNotIn("git", record)
        self.assertIn("processing", record["study_sections"])
        date = re.compile(r"\b20\d\d-\d\d-\d\d\b")
        self.assertIsNone(date.search(text))
        self.assertIsNone(date.search(changes))
        self.assertIn("## Build 1", changes)


class ChangeReportTests(unittest.TestCase):
    def test_a_new_exclusion_names_every_dependent_asset_with_rows_and_numbers(self):
        with study_folder(assay(), assay_study()) as path:
            build(path)
            edit(path, lambda study: study.update(processing=[EXCLUSION]))
            built = build(path)
            text = (built.out / "CHANGES.md").read_text(encoding="utf-8")
            record = json.loads((built.out / "provenance.json").read_text())
        entry = text.split("## Build 1")[0]
        self.assertIn("## Build 2", entry)
        self.assertIn("- study.yml changed in processing.", entry)
        self.assertIn(
            "- Step 3 is new: Exclude rows where Viability (%) is below 0. Reason: Instrument error. "
            "It removes 1 row.",
            entry,
        )
        self.assertIn("- Rows used: 36 → 35.", entry)
        for ident, asset in record["assets"].items():
            self.assertIn("processing", asset["depends_on"])
            line = next(line for line in entry.splitlines() if line.startswith(f"- **{ident}**"))
            self.assertIn("changed", line)
            self.assertIn("study.yml changed in processing;", line)
            # A figure counts the rows its analysis uses, so only the exclusion is the same for every asset.
            self.assertIn("(excluded 0 → 1, by step 3)", line)
            if asset["kind"] != "figure":
                self.assertIn("rows used 36 → 35 (excluded 0 → 1, by step 3)", line)
        self.assertIn("items to check 1 → 0", entry)
        self.assertIn("**descriptives** changed (Table 1)", entry)
        self.assertNotIn("Unchanged", entry)

    def test_assets_that_do_not_depend_on_a_change_are_named_unchanged(self):
        with study_folder(assay(), assay_study()) as path:
            build(path)
            edit(path, lambda study: study.update(sharing={"withhold": ["Passage"]}))
            built = build(path)
            entry = (built.out / "CHANGES.md").read_text(encoding="utf-8").split("## Build 1")[0]
        self.assertIn("- study.yml changed in sharing.", entry)
        self.assertIn("**processing-record** changed: study.yml changed in sharing; files changed:", entry)
        self.assertIn("Unchanged: data-audit, descriptives (Table 1)", entry)
        self.assertNotIn("**data-audit**", entry)

    def test_an_asset_that_depends_on_a_change_but_came_out_the_same_says_so(self):
        with study_folder(assay(), assay_study()) as path:
            build(path)
            edit(path, lambda study: study.update(decisions={"alpha": 0.01}))
            built = build(path)
            entry = (built.out / "CHANGES.md").read_text(encoding="utf-8").split("## Build 1")[0]
        self.assertEqual(built.changes, [])
        self.assertIn("- study.yml changed in decisions.", entry)
        self.assertIn("Although study.yml changed in decisions, data-audit", entry)
        self.assertIn("came out the same.", entry)

    def test_an_unchanged_rebuild_adds_no_entry(self):
        with study_folder(assay(), assay_study()) as path:
            first = build(path)
            text = (first.out / "CHANGES.md").read_text(encoding="utf-8")
            again = build(path)
            self.assertEqual((again.out / "CHANGES.md").read_text(encoding="utf-8"), text)

    def test_headline_numbers_show_enough_digits_to_differ(self):
        def record(numbers, rows, sha="a" * 64):
            return {
                "ada": "1",
                "inputs": [{"path": "data.csv", "sha256": sha}],
                "study_sections": {"analyses.primary": "1"},
                "assets": {
                    "results": {
                        "number": "Table 2",
                        "files": {"tables/results.csv": sha},
                        "depends_on": ["analyses.primary"],
                        "config_sha256": "same",
                        "rows": rows,
                        "numbers": numbers,
                    }
                },
            }

        before = record({"estimate": 1.23456, "p": 0.0412, "n": 36}, {"used": 36})
        after = record({"estimate": 1.23457, "p": 0.0412, "n": 35}, {"used": 35}, sha="b" * 64)
        (change,) = compare(before, after)
        self.assertEqual(change.number, "Table 2")
        self.assertIn("the input data changed", change.reasons)
        self.assertEqual(
            general_causes(before, after),
            ["the input data changed (data.csv sha256 aaaaaaaaaaaa… → bbbbbbbbbbbb…)"],
        )
        self.assertIn("rows used 36 → 35", change.reasons)
        self.assertIn("estimate 1.23456 → 1.23457", change.reasons)
        self.assertIn("n 36 → 35", change.reasons)
        self.assertFalse(any(reason.startswith("p ") for reason in change.reasons))

    def test_a_changed_rule_is_described_as_a_change_not_a_new_step(self):
        old = [
            {
                "step": 3,
                "kind": "exclude",
                "description": "Exclude rows where x is below 0",
                "reason": "r",
                "rows_before": 10,
                "rows_after": 9,
                "values_changed": 0,
            }
        ]
        new = [dict(old[0], description="Exclude rows where x is below 5", rows_after=6)]
        self.assertEqual(
            step_changes(old, new),
            [
                "Step 3 changed from “Exclude rows where x is below 0” to “Exclude rows where x is below 5”: "
                "it now removes 4 rows, where it removed 1 row."
            ],
        )
        self.assertEqual(step_changes(old, []), ["Step 3 was removed: Exclude rows where x is below 0."])


if __name__ == "__main__":
    unittest.main()
