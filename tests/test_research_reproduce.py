"""The reproducibility guarantee: a fresh copy of study.yml and the data rebuilds every asset.

These tests play two researchers. One builds the assets in folder A; the other
gets only what ``git clone`` would give (study.yml, the data and the custom
script) in folder B, rebuilds, and checks with ``ada verify``.
"""

import contextlib
import hashlib
import io
import json
import re
import shutil
import tempfile
import unittest
from pathlib import Path

import yaml
from research_helpers import assay, assay_study

from ada.cli import main
from ada.research import build
from ada.research.verify import Report, compare_table, compare_word, verify

EXCLUSION = {"exclude": {"where": {"column": "Viability (%)", "below": 0}}, "reason": "Instrument error."}

SCRIPT = """
from ada import research


@research.table("passage-means", title="Viability by passage group")
def passage_means(data, study):
    groups = (data["Passage"] >= 7).map({True: "7 or more", False: "under 7"})
    frame = data.groupby(groups)["Viability (%)"].agg(["count", "mean", "std"])
    return frame.rename_axis("Passage").reset_index()
"""


def write_study(root: Path, *, processing=None, custom: bool = True, sha: bool = False) -> Path:
    """The assay study in ``root``: data under data/, a custom script under analysis/, and study.yml."""
    (root / "data").mkdir(parents=True, exist_ok=True)
    assay().to_csv(root / "data" / "assay.csv", index=False)
    study = assay_study(processing=[EXCLUSION] if processing is None else processing)
    study["data"]["path"] = "data/assay.csv"
    if sha:
        study["data"]["sha256"] = hashlib.sha256((root / "data" / "assay.csv").read_bytes()).hexdigest()
    if custom:
        (root / "analysis").mkdir(exist_ok=True)
        (root / "analysis" / "extra.py").write_text(SCRIPT, encoding="utf-8")
        study["custom"] = ["analysis/extra.py"]
    path = root / "study.yml"
    path.write_text(yaml.safe_dump(study, sort_keys=False, allow_unicode=True), encoding="utf-8")
    return path


def clone(source: Path, target: Path) -> Path:
    """What another researcher's checkout holds: study.yml, the data and the custom scripts, nothing built."""
    study = yaml.safe_load((source / "study.yml").read_text(encoding="utf-8"))
    for relative in ["study.yml", study["data"]["path"], *study.get("custom", [])]:
        (target / relative).parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(source / relative, target / relative)
    return target / "study.yml"


def edit_study(path: Path, change) -> None:
    study = yaml.safe_load(path.read_text(encoding="utf-8"))
    change(study)
    path.write_text(yaml.safe_dump(study, sort_keys=False, allow_unicode=True), encoding="utf-8")


def rerecord(out: Path, relative: str) -> None:
    """Record a file's current checksum, as if the build had written it so: another machine's build."""
    record = json.loads((out / "provenance.json").read_text(encoding="utf-8"))
    digest = hashlib.sha256((out / relative).read_bytes()).hexdigest()
    for entry in record["assets"].values():
        if relative in entry["files"]:
            entry["files"][relative] = digest
    (out / "provenance.json").write_text(json.dumps(record, indent=1), encoding="utf-8")


def section(report, key: str) -> str:
    return "\n".join(line.text + "\n" + "\n".join(line.detail) for line in report.lines[key])


class FreshCheckoutTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.root = Path(self.folder.name)
        self.a, self.b = self.root / "a", self.root / "b"
        self.a.mkdir()
        self.b.mkdir()

    def tearDown(self):
        self.folder.cleanup()

    def test_a_fresh_checkout_rebuilds_every_asset_and_tracks_a_changed_exclusion(self):
        built = build(write_study(self.a))
        assets = json.loads((built.out / "provenance.json").read_text())["assets"]
        self.assertIn("passage-means", assets)

        path = clone(self.a, self.b)
        self.assertFalse((self.b / "paper-assets").exists())
        rebuilt = build(path)
        report = verify(path, out=built.out)
        self.assertTrue(report.ok, report.text())
        self.assertIn("Everything matches", report.summary)
        self.assertIn("analysis/extra.py", section(report, "rebuild"))
        for relative in assets["descriptives"]["files"]:
            self.assertEqual(
                (built.out / relative).read_bytes(), (rebuilt.out / relative).read_bytes(), relative
            )

        edit_study(path, lambda study: study["processing"][0]["exclude"]["where"].update(below=85))
        report = verify(path)
        self.assertFalse(report.ok)
        stale = section(report, "stale")
        dependent = [ident for ident, entry in assets.items() if "processing" in entry["depends_on"]]
        self.assertGreaterEqual(len(dependent), 5)
        for ident in dependent:
            self.assertIn(ident, stale)
        self.assertIn("study.yml changed in processing", stale)

        before = len(rebuilt.data)
        after = build(path)
        text = (after.out / "CHANGES.md").read_text(encoding="utf-8")
        self.assertLess(len(after.data), before)
        for ident in dependent:
            self.assertIn(f"**{ident}** changed", text)
        self.assertIn(f"rows used {before} → {len(after.data)}", text)
        self.assertIn("Step 3 changed from", text)
        self.assertTrue(verify(path).ok)


class VerifyTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.root = Path(self.folder.name)

    def tearDown(self):
        self.folder.cleanup()

    def built(self, **options):
        path = write_study(self.root, **options)
        return path, build(path).out

    def test_a_changed_data_file_is_caught(self):
        path, out = self.built(custom=False)
        frame = assay()
        frame.loc[0, "Viability (%)"] = 50.0
        frame.to_csv(self.root / "data" / "assay.csv", index=False)
        report = verify(path)
        self.assertFalse(report.ok)
        self.assertIn("not the file the assets were built from", section(report, "inputs"))
        self.assertIn("tables/descriptives.csv", section(report, "rebuild"))

    def test_a_data_file_that_differs_from_study_yml_stops_the_rebuild(self):
        path, out = self.built(custom=False, sha=True)
        with open(self.root / "data" / "assay.csv", "a", encoding="utf-8") as handle:
            handle.write("Z99,Control,5,88.0\n")
        report = verify(path)
        self.assertFalse(report.ok)
        self.assertIn("data/assay.csv has changed", section(report, "inputs"))
        self.assertIn("data.sha256", section(report, "inputs"))
        self.assertIn("Not rebuilt", section(report, "rebuild"))

    def test_a_hand_edited_table_is_caught_with_its_cell(self):
        path, out = self.built(custom=False)
        table = out / "tables" / "descriptives.csv"
        lines = table.read_text(encoding="utf-8").splitlines()
        cells = lines[1].split(",")
        mean = lines[0].split(",").index("mean")
        old = cells[mean]
        cells[mean] = "99.5"
        lines[1] = ",".join(cells)
        table.write_text("\n".join(lines) + "\n", encoding="utf-8")
        report = verify(path)
        self.assertFalse(report.ok)
        self.assertIn("Changed since the build: tables/descriptives.csv", section(report, "edited"))
        rebuild = section(report, "rebuild")
        self.assertIn("line 2 (Viability (%), Overall), column mean", rebuild)
        self.assertIn(f"99.5 on disk, {old} rebuilt", rebuild)

    def test_a_deleted_file_is_reported(self):
        path, out = self.built(custom=False)
        (out / "tables" / "descriptives.tex").unlink()
        report = verify(path)
        self.assertIn("Deleted since the build: tables/descriptives.tex", section(report, "edited"))
        self.assertNotIn("descriptives.tex", section(report, "rebuild"))

    def test_a_missing_data_file_gives_a_clear_message(self):
        path, out = self.built(custom=False)
        (self.root / "data" / "assay.csv").unlink()
        report = verify(path)
        self.assertFalse(report.ok)
        self.assertIn("The data file data/assay.csv is not there", section(report, "inputs"))
        self.assertIn("Not rebuilt", section(report, "rebuild"))

    def test_a_study_never_built_says_to_build_first(self):
        path = write_study(self.root, custom=False)
        report = verify(path)
        self.assertFalse(report.ok)
        self.assertIn("Run `ada build` first", section(report, "inputs"))
        self.assertIn("nothing has been built", report.summary)

    def test_tiny_float_noise_passes_and_a_real_change_fails(self):
        path, out = self.built(custom=False)
        table = out / "tables" / "descriptives.csv"
        original = table.read_text(encoding="utf-8")
        header, first = original.splitlines()[:2]
        mean = header.split(",").index("mean")
        value = float(first.split(",")[mean])

        def replace(new: float) -> None:
            cells = first.split(",")
            cells[mean] = repr(new)
            table.write_text(original.replace(first, ",".join(cells), 1), encoding="utf-8")
            rerecord(out, "tables/descriptives.csv")

        replace(value * (1 + 1e-12))
        self.assertTrue(verify(path).ok, verify(path).text())
        replace(value + 0.01)
        report = verify(path)
        self.assertFalse(report.ok)
        self.assertIn("column mean", section(report, "rebuild"))
        self.assertTrue(verify(path, rtol=1e-3).ok)

    def test_a_changed_custom_script_makes_its_assets_stale_and_a_new_table_missing(self):
        path, out = self.built()
        script = self.root / "analysis" / "extra.py"
        script.write_text(
            script.read_text().replace('"std"', '"median"') + '\n\n@research.table("wells", title="Wells")\n'
            "def wells(data, study):\n"
            '    return data.groupby("Treatment", observed=True).size().rename("n").reset_index()\n',
            encoding="utf-8",
        )
        report = verify(path)
        stale = section(report, "stale")
        self.assertIn("Out of date because analysis/extra.py changed since the build: passage-means", stale)
        self.assertIn("Missing: study.yml now asks for wells", stale)
        self.assertIn("as expected while their assets are out of date", section(report, "rebuild"))

    def test_a_different_ada_version_is_reported(self):
        path, out = self.built(custom=False)
        record = json.loads((out / "provenance.json").read_text())
        record["ada"] = "0.0.1"
        (out / "provenance.json").write_text(json.dumps(record), encoding="utf-8")
        report = verify(path)
        self.assertFalse(report.ok)
        self.assertIn("built by ADA 0.0.1", section(report, "stale"))

    def test_an_image_that_differs_only_in_bytes_is_a_note(self):
        path, out = self.built()
        script = self.root / "analysis" / "extra.py"
        script.write_text(
            script.read_text() + "\n\n@research.figure('spread', title='Spread')\n"
            "def spread(data, study):\n"
            "    import matplotlib.pyplot as plt\n"
            "    fig, ax = plt.subplots(figsize=research.figure_size(study))\n"
            "    ax.plot(data['Passage'], data['Viability (%)'], 'o')\n"
            "    return fig, data[['Passage', 'Viability (%)']]\n",
            encoding="utf-8",
        )
        out = build(path).out
        image = out / "figures" / "spread.png"
        image.write_bytes(image.read_bytes() + b"\0")
        rerecord(out, "figures/spread.png")
        report = verify(path)
        self.assertTrue(report.ok, report.text())
        self.assertIn("figures/spread.png", section(report, "rebuild"))
        self.assertIn("source data match", section(report, "rebuild"))

    def test_the_command_exits_0_when_everything_matches_and_1_otherwise(self):
        path, out = self.built(custom=False)
        printed = io.StringIO()
        with contextlib.redirect_stdout(printed):
            self.assertEqual(main(["verify", "--study", str(path), "--rtol", "1e-9", "--atol", "1e-12"]), 0)
        self.assertIn("Everything matches", printed.getvalue())
        (out / "tables" / "descriptives.tex").write_text("edited\n", encoding="utf-8")
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(main(["verify", "--study", str(path)]), 1)

    def test_the_json_flag_prints_the_same_findings_and_exits_as_the_report(self):
        path, out = self.built(custom=False)
        printed = io.StringIO()
        with contextlib.redirect_stdout(printed):
            self.assertEqual(main(["verify", "--study", str(path), "--json"]), 0)
        payload = json.loads(printed.getvalue())
        self.assertEqual(list(payload), ["study", "ok", "summary", "inputs", "stale", "edited", "rebuild"])
        self.assertTrue(payload["ok"])
        self.assertTrue(payload["summary"].startswith("Everything matches"), payload["summary"])
        for key in ("inputs", "stale", "edited"):
            self.assertNotIn("fail", [line["status"] for line in payload[key]])

        (out / "tables" / "descriptives.tex").write_text("edited\n", encoding="utf-8")
        as_text, as_json = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(as_text):
            text_code = main(["verify", "--study", str(path)])
        with contextlib.redirect_stdout(as_json):
            self.assertEqual(main(["verify", "--study", str(path), "--json"]), text_code)
        payload = json.loads(as_json.getvalue())
        self.assertFalse(payload["ok"])
        edited = [line for line in payload["edited"] if line["status"] == "fail"]
        self.assertEqual(len(edited), 1)
        self.assertIn("Changed since the build: tables/descriptives.tex", edited[0]["text"])
        self.assertIn("Changed since the build", as_text.getvalue())

    def test_the_json_of_a_study_never_built_says_so_under_inputs(self):
        write_study(self.root, custom=False)
        printed = io.StringIO()
        with contextlib.chdir(self.root), contextlib.redirect_stdout(printed):
            self.assertEqual(main(["verify", "--json"]), 1)
        payload = json.loads(printed.getvalue())
        self.assertFalse(payload["ok"])
        self.assertIn("nothing has been built", payload["summary"])
        last = payload["inputs"][-1]
        self.assertEqual(last["status"], "fail")
        self.assertIn("There is no build in paper-assets to verify", last["text"])
        self.assertIn("Run `ada build` first", last["text"])
        self.assertEqual((payload["stale"], payload["edited"], payload["rebuild"]), ([], [], []))

    def test_the_json_holds_no_absolute_paths_or_dates_and_is_the_same_twice(self):
        _, out = self.built(custom=False)
        runs = []
        with contextlib.chdir(self.root):
            for _ in range(2):
                printed = io.StringIO()
                with contextlib.redirect_stdout(printed):
                    self.assertEqual(main(["verify", "--json"]), 0)
                runs.append(printed.getvalue())
        self.assertEqual(runs[0], runs[1])
        payload = json.loads(runs[0])
        self.assertEqual(payload["study"], "study.yml")
        self.assertNotIn(str(self.root), runs[0])
        self.assertNotIn(str(out.parent), runs[0])
        self.assertIsNone(re.search(r"\b20\d\d-\d\d-\d\d\b", runs[0]))


class ReportTests(unittest.TestCase):
    def test_the_json_and_the_text_come_from_the_same_findings(self):
        report = Report(study="study.yml")
        report.add("inputs", "ok", "data/assay.csv is there.", ["sha256 3f2a9c81d0e4…"])
        report.add("edited", "fail", "Changed since the build: tables/descriptives.csv.")
        report.summary = "Does not match: see Edited since the build (1) above."
        payload = report.to_dict()
        self.assertEqual(payload["study"], "study.yml")
        self.assertFalse(payload["ok"])
        self.assertEqual(payload["summary"], report.summary)
        self.assertEqual(
            payload["inputs"],
            [{"status": "ok", "text": "data/assay.csv is there.", "detail": ["sha256 3f2a9c81d0e4…"]}],
        )
        self.assertEqual([line["text"] for line in payload["edited"]], [report.lines["edited"][0].text])
        self.assertEqual((payload["stale"], payload["rebuild"]), ([], []))
        text = report.text()
        for shown in ("Inputs", "data/assay.csv is there.", "Edited since the build"):
            self.assertIn(shown, text)
        self.assertIn("Changed since the build", text)
        self.assertTrue(json.loads(json.dumps(payload)))


class ComparisonTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.root = Path(self.folder.name)

    def tearDown(self):
        self.folder.cleanup()

    def write(self, name: str, text: str) -> Path:
        path = self.root / name
        path.write_text(text, encoding="utf-8")
        return path

    def test_numbers_match_within_the_tolerance_and_text_exactly(self):
        disk = self.write("a.csv", "group,mean,label\nA,1.0000000000001,x\nB,2.5,y\n")
        close = self.write("b.csv", "group,mean,label\nA,1.0,x\nB,2.5,y\n")
        self.assertEqual(compare_table(disk, close), ([], 0))
        apart = self.write("c.csv", "group,mean,label\nA,1.001,x\nB,2.5,z\n")
        differences, total = compare_table(disk, apart)
        self.assertEqual(total, 2)
        self.assertEqual(differences[0], "line 2 (A), column mean: 1.0000000000001 on disk, 1.001 rebuilt")
        self.assertEqual(differences[1], "line 3 (B), column label: y on disk, z rebuilt")
        self.assertEqual(compare_table(disk, apart, rtol=1e-2)[1], 1)

    def test_differences_beyond_the_cap_are_counted(self):
        rows = "\n".join(f"{i},{i}.5" for i in range(20))
        disk = self.write("a.csv", "id,value\n" + rows + "\n")
        shifted = self.write("b.csv", "id,value\n" + rows.replace(".5", ".7") + "\n")
        differences, total = compare_table(disk, shifted)
        self.assertEqual((len(differences), total), (5, 20))

    def test_columns_and_rows_must_match(self):
        disk = self.write("a.csv", "x,y\n1,2\n3,4\n")
        self.assertIn("columns on disk: x, y", compare_table(disk, self.write("b.csv", "x,z\n1,2\n3,4\n"))[0])
        self.assertIn("2 rows on disk, 1 rebuilt", compare_table(disk, self.write("c.csv", "x,y\n1,2\n"))[0])

    def test_word_files_are_compared_by_their_text(self):
        from docx import Document

        def word(name: str, text: str) -> Path:
            document = Document()
            document.add_paragraph(text)
            document.add_table(rows=1, cols=1).rows[0].cells[0].text = "12.3"
            document.save(self.root / name)
            return self.root / name

        first = word("a.docx", "Table 1. Viability")
        same = word("b.docx", "Table 1. Viability")
        self.assertEqual(compare_word(first, same), ([], 0))
        differences, total = compare_word(first, word("c.docx", "Table 1. Survival"))
        self.assertEqual(total, 1)
        self.assertIn('paragraph 1: "Table 1. Viability" on disk, "Table 1. Survival" rebuilt', differences)


if __name__ == "__main__":
    unittest.main()
